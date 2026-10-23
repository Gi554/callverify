"""Shared tabular exports (CSV for any tool, XLSX for Excel), safe against spreadsheet formulas."""

import csv
import re
from datetime import date, datetime
from io import BytesIO, StringIO
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError

# Excel generation is much slower than CSV: larger volumes go through CSV.
MAX_ROWS = {'csv': 50000, 'xlsx': 20000}
KINDS = ('csv', 'xlsx')
CONTROL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
# Phone-like values may start with + or - without being formulas.
NUMBERISH = re.compile(r'^[+-]?[\d\s().-]*$')


def kind(request):
    value = request.query_params.get('file', '')
    if value not in KINDS:
        raise ValidationError('Format d’export attendu : csv ou xlsx.')
    if 'page' in request.query_params:
        raise ValidationError('L’export porte sur tous les résultats filtrés, sans pagination.')
    return value


def bounded(queryset, kind):
    # A single bounded read avoids a count/read race and never silently truncates.
    limit = MAX_ROWS[kind]
    rows = list(queryset[: limit + 1])
    if len(rows) > limit:
        hint = ' ou choisissez CSV' if kind == 'xlsx' else ''
        raise ValidationError(
            f'Plus de {limit} lignes : précisez les filtres{hint} avant d’exporter.'
        )
    return rows


def plain(value):
    if isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None) if timezone.is_aware(value) else value
    if isinstance(value, bool):
        return 'Oui' if value else 'Non'
    if isinstance(value, str):
        return CONTROL.sub('�', value)
    return value


def csv_text(value):
    value = plain(value)
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%d/%m/%Y %H:%M')
    if isinstance(value, date):
        return value.strftime('%d/%m/%Y')
    value = str(value)
    # Neutralise formula triggers (CSV injection) while keeping phone numbers readable.
    if value[:1] in ('=', '@', '\t', '\r') or (value[:1] in '+-' and not NUMBERISH.match(value)):
        return "'" + value
    return value


def response(kind, headers, rows, filename, title):
    stamp = timezone.localdate().isoformat()
    if kind == 'csv':
        stream = StringIO()
        # Semicolons and a BOM let French-locale Excel open the file directly.
        writer = csv.writer(stream, delimiter=';', lineterminator='\r\n')
        writer.writerow(headers)
        for row in rows:
            writer.writerow([csv_text(value) for value in row])
        result = HttpResponse('﻿' + stream.getvalue(), content_type='text/csv; charset=utf-8')
    else:
        book = Workbook(write_only=True)
        sheet = book.create_sheet(title[:31])
        sheet.freeze_panes = 'A2'
        for index in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 22
        head = []
        for header in headers:
            cell = WriteOnlyCell(sheet, header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='53661C')
            head.append(cell)
        sheet.append(head)
        for row in rows:
            values = []
            for value in row:
                value = plain(value)
                if isinstance(value, str) and value.startswith('='):
                    # Explicit text cells keep literal values instead of formulas.
                    cell = WriteOnlyCell(sheet, value)
                    cell.data_type = 's'
                    value = cell
                elif isinstance(value, datetime):
                    cell = WriteOnlyCell(sheet, value)
                    cell.number_format = 'dd/mm/yyyy hh:mm'
                    value = cell
                values.append(value)
            sheet.append(values)
        stream = BytesIO()
        book.save(stream)
        result = HttpResponse(
            stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
    result['Content-Disposition'] = f'attachment; filename="{filename}-{stamp}.{kind}"'
    result['Cache-Control'] = 'no-store'
    result['X-Content-Type-Options'] = 'nosniff'
    return result
