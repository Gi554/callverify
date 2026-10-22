from django.test import TestCase,override_settings
from rest_framework.test import APIClient
from packages.access.tickets import issue
from django.conf import settings
from unittest import skipUnless
if 'apps.assurance.backend.crm' in settings.INSTALLED_APPS:
    from .models import Contact

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class AssuranceTests(TestCase):
    def setUp(self):self.client=APIClient()
    def authorize(self,subject,role,space='assurance'):
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(subject,space,role))
    def test_isolation_ticket_and_anonymous_are_denied(self):
        self.assertEqual(self.client.get('/api/assurance/contacts/').status_code,401)
        self.authorize(1,'administrator','isolation')
        self.assertEqual(self.client.get('/api/assurance/contacts/').status_code,401)
    def test_advisor_only_sees_assigned_contact_and_no_setter_fields(self):
        visible=Contact.objects.create(name='A',phone='original phone',assigned_advisor_id=7,created_by_id=123)
        hidden=Contact.objects.create(name='B',phone='other',assigned_advisor_id=8,created_by_id=124)
        self.authorize(7,'advisor')
        rows=self.client.get('/api/assurance/contacts/').data['results']
        self.assertEqual([r['id'] for r in rows],[visible.pk])
        self.assertNotIn('created_by_id',rows[0]);self.assertNotIn('assigned_advisor_id',rows[0])
        self.assertEqual(self.client.get(f'/api/assurance/contacts/{hidden.pk}/').status_code,404)
    def test_contact_read_preserves_original_phone_without_technical_processing(self):
        Contact.objects.create(name='Client',phone='Original tel 01.00.00',created_by_id=1)
        self.authorize(1,'administrator')
        response=self.client.get('/api/assurance/contacts/')
        self.assertEqual(response.data['results'][0]['phone'],'Original tel 01.00.00')
    def test_setter_edit_is_blocked_until_locking_exists(self):
        self.authorize(1,'setter')
        self.assertEqual(self.client.post('/api/assurance/contacts/',{'name':'A','phone':'x'}).status_code,405)
    def test_no_twilio_or_normalization_endpoints(self):
        self.authorize(1,'administrator')
        self.assertEqual(self.client.post('/api/assurance/lookup/').status_code,404)
        self.assertEqual(self.client.post('/api/assurance/calls/').status_code,404)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class RawImportTests(TestCase):
    def setUp(self):
        self.client=APIClient()
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(17,'assurance','administrator'))
        self.url='/api/assurance/imports/'
    def upload(self,content=b'Nom;Telephone;Courriel\nAlice;001 abc;not-an-email\n',**kwargs):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return self.client.post(self.url,{'file':SimpleUploadedFile('contacts.csv',content),'has_header':'true',**kwargs},format='multipart')
    def confirm(self,content=None,**kwargs):
        from uuid import uuid4
        import json
        args={'mode':'confirm','request_id':str(uuid4()),'mapping':json.dumps({'name':0,'phone':1,'email':2}),**kwargs}
        return self.upload(**({'content':content} if content is not None else {}),**args)
    def test_preview_has_no_side_effects_and_preserves_strings(self):
        from .models import ContactImport
        response=self.upload()
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.data['rows'][0]['values'][1],'001 abc')
        self.assertEqual(Contact.objects.count(),0);self.assertEqual(ContactImport.objects.count(),0)
    def test_raw_duplicates_and_bad_format_are_imported_without_lookup(self):
        content=b'Nom;Telephone;Courriel\nAlice; 001 abc ;bad-email\nAlice; 001 abc ;bad-email\n'
        response=self.confirm(content)
        self.assertEqual(response.status_code,201,response.data)
        self.assertEqual(response.data['imported'],2)
        self.assertEqual(list(Contact.objects.values_list('phone',flat=True)),[' 001 abc ',' 001 abc '])
        self.assertEqual(Contact.objects.first().source_values,['Alice',' 001 abc ','bad-email'])
        self.assertEqual(Contact.objects.first().import_batch.mapping,{'name':0,'phone':1,'email':2})
    def test_rejected_rows_report_physical_line_and_valid_rows_survive(self):
        response=self.confirm(b'Nom;Telephone;Courriel\n\nMissing;;\nOK;raw;\n')
        self.assertEqual(response.status_code,201,response.data)
        self.assertEqual(response.data['total'],2);self.assertEqual(response.data['imported'],1)
        self.assertEqual(response.data['report'],[{'line':3,'reason':'Téléphone manquant'}])
        self.assertEqual(Contact.objects.first().source_row,4)
    def test_idempotent_retry_and_conflicting_payload(self):
        from uuid import uuid4
        key=str(uuid4())
        first=self.confirm(request_id=key);retry=self.confirm(request_id=key)
        self.assertEqual(first.data['id'],retry.data['id']);self.assertEqual(Contact.objects.count(),1)
        conflict=self.confirm(b'Nom;Telephone;Courriel\nOther;changed;\n',request_id=key)
        self.assertEqual(conflict.status_code,409);self.assertEqual(Contact.objects.count(),1)
    def test_no_header_keeps_first_line(self):
        import json
        response=self.confirm(b'001\n002\n',has_header='false',mapping=json.dumps({'phone':0}))
        self.assertEqual(response.status_code,201,response.data)
        self.assertEqual(list(Contact.objects.values_list('phone',flat=True)),['001','002'])
    def test_xlsx_first_sheet_and_formula_are_not_evaluated(self):
        from openpyxl import Workbook
        from io import BytesIO
        from django.core.files.uploadedfile import SimpleUploadedFile
        from uuid import uuid4
        book=Workbook();sheet=book.active;sheet.append(['Phone']);sheet.append(['=1+1'])
        book.create_sheet('Ignore').append(['other'])
        stream=BytesIO();book.save(stream)
        response=self.client.post(self.url,{'file':SimpleUploadedFile('sample.xlsx',stream.getvalue()),'mode':'confirm','mapping':'{"phone":0}','request_id':str(uuid4())},format='multipart')
        self.assertEqual(response.status_code,201,response.data)
        self.assertEqual(Contact.objects.get().phone,'=1+1')
    def test_permissions_protect_imports_and_history(self):
        self.confirm()
        for role in ('setter','advisor'):
            self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(5,'assurance',role))
            self.assertEqual(self.upload().status_code,403)
            self.assertEqual(self.client.get(self.url).status_code,403)
    def test_invalid_mapping_does_not_write_contacts_or_batches(self):
        from .models import ContactImport
        for mapping in ('{}','{"phone":99}','{"phone":0,"name":0}','{"phone":true}','[]'):
            response=self.confirm(mapping=mapping)
            self.assertEqual(response.status_code,400,response.data)
        self.assertEqual(Contact.objects.count(),0);self.assertEqual(ContactImport.objects.count(),0)
    def test_bad_file_and_row_limit(self):
        from unittest.mock import patch
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.assertEqual(self.client.post(self.url,{'file':SimpleUploadedFile('bad.xlsx',b'bad')},format='multipart').status_code,400)
        with patch('apps.assurance.backend.crm.imports.MAX_ROWS',1):
            self.assertEqual(self.upload(b'Phone\na\nb\n').status_code,400)
        with patch('apps.assurance.backend.crm.imports.MAX_BYTES',2):
            self.assertEqual(self.upload().status_code,400)
    def test_search_and_pagination_never_expand_advisor_scope(self):
        own=Contact.objects.create(name='Visible',phone='raw',created_by_id=17,assigned_advisor_id=5,source_values=['secret setter'])
        Contact.objects.create(name='Hidden',phone='raw',created_by_id=17,assigned_advisor_id=6)
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(5,'assurance','advisor'))
        response=self.client.get('/api/assurance/contacts/?search=raw')
        self.assertEqual(response.data['count'],1)
        self.assertEqual(response.data['results'][0]['id'],own.pk)
        self.assertNotIn('source_values',response.data['results'][0])
        self.assertNotIn('import_batch',response.data['results'][0])
        self.assertEqual(self.client.get('/api/assurance/contacts/?search=Hidden').data['count'],0)
    def test_import_history_excludes_actor_and_raw_data(self):
        self.confirm()
        response=self.client.get(self.url)
        self.assertEqual(response.data['count'],1)
        self.assertNotIn('created_by_id',response.data['results'][0])
        self.assertNotIn('source_values',str(response.data))
    def test_5000_rows_are_imported_without_deduplication(self):
        content=('Nom;Telephone;Courriel\n'+'Volume;unverified;raw\n'*5000).encode()
        response=self.confirm(content)
        self.assertEqual(response.status_code,201,response.data)
        self.assertEqual(response.data['imported'],5000)
        self.assertEqual(Contact.objects.count(),5000)
    def test_audit_failure_rolls_back_batch_and_contacts(self):
        from unittest.mock import patch
        from .models import ContactImport
        with patch('apps.assurance.backend.crm.imports.BusinessAudit.objects.create',side_effect=RuntimeError('audit unavailable')):
            with self.assertRaises(RuntimeError):self.confirm()
        self.assertFalse(Contact.objects.exists());self.assertFalse(ContactImport.objects.exists())

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class ContactLeaseTests(TestCase):
    def setUp(self):
        self.contact=Contact.objects.create(name='Original',phone='raw',email='unchanged',created_by_id=17,assigned_advisor_id=8)
        self.client=APIClient();self.authorize(1)
        self.url=f'/api/assurance/contacts/{self.contact.pk}/'
    def authorize(self,subject,role='setter'):
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(subject,'assurance',role))
    def claim(self):
        response=self.client.post(self.url+'lease/',{},format='json')
        self.assertEqual(response.status_code,201,response.data)
        return response.data['token']
    def save(self,token,**values):
        return self.client.patch(self.url,{'token':token,'revision':0,'name':'Changed',**values},format='json')
    def test_second_setter_reads_but_cannot_claim_or_edit(self):
        token=self.claim();self.authorize(2)
        self.assertEqual(self.client.get(self.url).status_code,200)
        self.assertEqual(self.client.post(self.url+'lease/',{},format='json').status_code,409)
        self.assertEqual(self.save(token).status_code,409)
        self.contact.refresh_from_db();self.assertEqual(self.contact.name,'Original')
    def test_renew_release_and_reclaim(self):
        token=self.claim()
        self.assertEqual(self.client.patch(self.url+'lease/',{'token':token},format='json').status_code,200)
        self.assertEqual(self.client.delete(self.url+'lease/',{'token':token},format='json').status_code,204)
        fresh=self.claim();self.assertNotEqual(fresh,token)
        self.assertEqual(self.save(token).status_code,409)
        self.assertEqual(self.save(fresh).status_code,200)
    def test_expiry_denies_save_renew_and_allows_another_setter(self):
        from django.utils import timezone
        from datetime import timedelta
        token=self.claim()
        Contact.objects.filter(pk=self.contact.pk).update(lease_expires_at=timezone.now()-timedelta(seconds=1))
        self.assertEqual(self.save(token).status_code,409)
        self.assertEqual(self.client.patch(self.url+'lease/',{'token':token},format='json').status_code,409)
        self.authorize(2);self.claim()
    def test_save_preserves_raw_values_and_appends_audit_with_revision(self):
        from .models import BusinessAudit
        token=self.claim()
        response=self.save(token,phone=' 001 weird ',email='not-validated')
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.data['phone'],' 001 weird ');self.assertEqual(response.data['revision'],1)
        entry=BusinessAudit.objects.get(action='contact.updated')
        self.assertEqual(entry.actor_id,1);self.assertEqual(entry.metadata['before']['name'],'Original')
        self.assertEqual(self.save(token,name='Stale').status_code,409)
        self.assertEqual(self.save(token,revision=1,name='Next').status_code,200)
        self.assertEqual(BusinessAudit.objects.filter(action='contact.updated').count(),2)
    def test_edit_requires_valid_token_and_rejects_protected_fields(self):
        token=self.claim()
        self.assertEqual(self.save('bad').status_code,400)
        self.assertEqual(self.save(token,phone='   ').status_code,400)
        for field in ('created_by_id','assigned_advisor_id','status','source_values','lease_owner_id'):
            self.assertEqual(self.save(token,**{field:99}).status_code,400)
    def test_admin_override_is_justified_audited_and_fences_old_token(self):
        from .models import BusinessAudit
        token=self.claim()
        self.assertEqual(self.client.post(self.url+'unlock/',{'reason':'test'},format='json').status_code,403)
        self.authorize(10,'administrator')
        self.assertEqual(self.client.post(self.url+'unlock/',{},format='json').status_code,400)
        self.assertEqual(self.client.post(self.url+'unlock/',{'reason':'Reprise après absence'},format='json').status_code,204)
        fresh=self.claim();self.assertNotEqual(fresh,token)
        self.authorize(1);self.assertEqual(self.save(token).status_code,409)
        entry=BusinessAudit.objects.get(action='contact.force_unlocked')
        self.assertEqual(entry.metadata['reason'],'Reprise après absence')
    def test_advisor_never_sees_setter_or_lease_metadata_and_cannot_edit(self):
        self.claim();self.authorize(8,'advisor')
        response=self.client.get(self.url)
        self.assertEqual(response.status_code,200)
        self.assertEqual(set(response.data),{'id','name','phone','email','status','status_label','updated_at','revision'})
        for suffix in ('lease/','history/'):
            self.assertEqual(self.client.get(self.url+suffix).status_code,403)
        self.assertEqual(self.client.patch(self.url,{},format='json').status_code,403)
        self.assertEqual(self.client.post(self.url+'lease/',{},format='json').status_code,403)
    def test_history_is_admin_only_and_scoped_to_contact(self):
        from .models import BusinessAudit
        token=self.claim();self.save(token)
        BusinessAudit.objects.create(actor_id=1,object_id=self.contact.pk+100,action='contact.updated')
        self.assertEqual(self.client.get(self.url+'history/').status_code,403)
        self.authorize(10,'administrator');response=self.client.get(self.url+'history/')
        self.assertEqual(response.data['count'],2)
        self.assertEqual(response.data['results'][0]['action'],'contact.updated')
    def test_audit_failure_rolls_back_the_edit(self):
        from unittest.mock import patch
        token=self.claim()
        with patch('apps.assurance.backend.crm.leases.BusinessAudit.objects.create',side_effect=RuntimeError('audit failure')):
            with self.assertRaises(RuntimeError):self.save(token)
        self.contact.refresh_from_db();self.assertEqual(self.contact.revision,0);self.assertEqual(self.contact.name,'Original')

from django.test import TransactionTestCase
@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class ConcurrentLeaseTests(TransactionTestCase):
    def test_simultaneous_claims_have_only_one_winner_on_postgres(self):
        from django.db import connection,connections
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        if connection.vendor!='postgresql':self.skipTest('Concurrence réelle vérifiée sur PostgreSQL dans la CI dédiée')
        contact=Contact.objects.create(name='Concurrent',phone='raw',created_by_id=17)
        barrier=Barrier(2)
        def claim(subject):
            try:
                client=APIClient();client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(subject,'assurance','setter'))
                barrier.wait(timeout=10)
                return client.post(f'/api/assurance/contacts/{contact.pk}/lease/',{},format='json').status_code
            finally:connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(claim,[41,42]))
        self.assertEqual(sorted(results),[201,409])
    def test_simultaneous_saves_do_not_overwrite_each_other_on_postgres(self):
        from django.db import connection,connections
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        if connection.vendor!='postgresql':self.skipTest('Concurrence réelle vérifiée sur PostgreSQL dans la CI dédiée')
        contact=Contact.objects.create(name='Concurrent',phone='raw',created_by_id=17)
        client=APIClient();client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(41,'assurance','setter'))
        token=client.post(f'/api/assurance/contacts/{contact.pk}/lease/',{},format='json').data['token']
        barrier=Barrier(2)
        def save(name):
            try:
                api=APIClient();api.credentials(HTTP_AUTHORIZATION='Bearer '+issue(41,'assurance','setter'))
                barrier.wait(timeout=10)
                return api.patch(f'/api/assurance/contacts/{contact.pk}/',{'name':name,'revision':0,'token':token},format='json').status_code
            finally:connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(save,['First','Second']))
        self.assertEqual(sorted(results),[200,409])
        contact.refresh_from_db();self.assertEqual(contact.revision,1)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class AssignmentTests(TestCase):
    def setUp(self):
        self.client=APIClient();self.authorize(10,'administrator');self.url='/api/assurance/assignments/'
    def authorize(self,subject,role):self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(subject,'assurance',role))
    def proof(self,setter=21,advisor=31,actor=10):
        from packages.access.assignments import issue_assignment
        return issue_assignment(actor,setter,'Setter '+str(setter),advisor,'Advisor '+str(advisor))
    def create(self,**kwargs):return self.client.post(self.url,{'proof':self.proof(**kwargs)},format='json')
    def test_multiple_advisors_per_setter_and_duplicate_pair_conflict(self):
        from .models import AdvisorAssignment
        self.assertEqual(self.create().status_code,201)
        self.assertEqual(self.create(advisor=32).status_code,201)
        self.assertEqual(self.create().status_code,409)
        self.assertEqual(AdvisorAssignment.objects.count(),2)
    def test_create_does_not_trust_unsigned_ids_or_wrong_actor(self):
        self.assertEqual(self.client.post(self.url,{'setter_id':21,'advisor_id':31},format='json').status_code,400)
        self.assertEqual(self.create(actor=11).status_code,400)
        self.assertEqual(self.client.post(self.url,{'proof':issue(10,'assurance','administrator')},format='json').status_code,400)
    def test_expired_or_tampered_attestation_is_rejected(self):
        from unittest.mock import patch
        with patch('django.core.signing.time.time',return_value=1000):expired=self.proof()
        for proof in (expired,self.proof()+'x'):
            self.assertEqual(self.client.post(self.url,{'proof':proof},format='json').status_code,400)
    def test_setter_only_reads_own_active_advisors_without_setter_metadata(self):
        first=self.create().data
        self.create(setter=22,advisor=32)
        suspended=self.create(advisor=33).data
        self.client.patch(self.url+str(suspended['id'])+'/',{'active':False,'revision':0},format='json')
        self.authorize(21,'setter');response=self.client.get(self.url)
        self.assertEqual(response.data['count'],1)
        self.assertEqual(response.data['results'],[{'id':first['id'],'advisor':{'id':31,'name':'Advisor 31'},'active':True}])
        self.assertEqual(self.client.get(self.url+str(suspended['id'])+'/').status_code,404)
        self.assertEqual(self.client.get(self.url+str(first['id'])+'/history/').status_code,403)
        self.assertEqual(self.create().status_code,403)
        self.assertEqual(self.client.patch(self.url+str(first['id'])+'/',{'active':False,'revision':0},format='json').status_code,403)
    def test_advisor_cannot_discover_setters_via_collection_detail_or_history(self):
        row=self.create().data;self.authorize(31,'advisor')
        for path in (self.url,self.url+str(row['id'])+'/',self.url+str(row['id'])+'/history/'):
            self.assertEqual(self.client.get(path).status_code,403)
        self.assertEqual(self.create().status_code,403)
    def test_reactivation_needs_fresh_proof_for_exact_pair_and_keeps_history(self):
        row=self.create().data;url=self.url+str(row['id'])+'/'
        self.assertEqual(self.client.patch(url,{'active':False,'revision':0},format='json').status_code,200)
        self.assertEqual(self.client.patch(url,{'active':True,'revision':1},format='json').status_code,400)
        self.assertEqual(self.client.patch(url,{'active':True,'revision':1,'proof':self.proof(advisor=99)},format='json').status_code,400)
        self.assertEqual(self.client.patch(url,{'active':True,'revision':1,'proof':self.proof()},format='json').status_code,200)
        history=self.client.get(url+'history/').data
        self.assertEqual(history['count'],3)
        self.assertEqual(history['results'][0]['metadata']['before']['active'],False)
    def test_stale_revision_does_not_overwrite_latest_state(self):
        row=self.create().data;url=self.url+str(row['id'])+'/'
        self.assertEqual(self.client.patch(url,{'active':False,'revision':0},format='json').status_code,200)
        self.assertEqual(self.client.patch(url,{'active':True,'revision':0,'proof':self.proof()},format='json').status_code,409)
        self.assertFalse(self.client.get(url).data['active'])
    def test_filter_and_paginated_admin_history_are_scoped(self):
        row=self.create().data;self.create(setter=22,advisor=32)
        self.assertEqual(self.client.get(self.url+'?search=Setter+21').data['count'],1)
        self.assertEqual(self.client.get(self.url+'?active=false').data['count'],0)
        self.assertEqual(self.client.get(self.url+str(row['id'])+'/history/').data['count'],1)
    def test_audit_failure_rolls_back_association(self):
        from unittest.mock import patch
        from .models import AdvisorAssignment
        with patch('apps.assurance.backend.crm.assignments.BusinessAudit.objects.create',side_effect=RuntimeError('audit failure')):
            with self.assertRaises(RuntimeError):self.create()
        self.assertFalse(AdvisorAssignment.objects.exists())
    def test_association_does_not_grant_access_to_unassigned_prospects(self):
        Contact.objects.create(name='Private prospect',phone='raw',created_by_id=21)
        self.assertEqual(self.create().status_code,201)
        self.authorize(31,'advisor')
        self.assertEqual(self.client.get('/api/assurance/contacts/').data['count'],0)
    def test_missing_or_wrong_workspace_ticket_is_denied(self):
        self.client.credentials();self.assertEqual(self.client.get(self.url).status_code,401)
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(10,'isolation','administrator'))
        self.assertEqual(self.client.get(self.url).status_code,401)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class ConcurrentAssignmentTests(TransactionTestCase):
    def race(self,operation):
        from django.db import connection,connections
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        if connection.vendor!='postgresql':self.skipTest('Concurrence réelle vérifiée dans la CI PostgreSQL')
        barrier=Barrier(2)
        def execute(_):
            try:
                client=APIClient();client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(10,'assurance','administrator'))
                barrier.wait(timeout=10)
                return operation(client).status_code
            finally:connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:return list(pool.map(execute,[0,1]))
    def test_duplicate_simultaneous_creates_have_one_winner(self):
        from packages.access.assignments import issue_assignment
        from .models import AdvisorAssignment
        proof=issue_assignment(10,21,'Setter',31,'Advisor')
        results=self.race(lambda client:client.post('/api/assurance/assignments/',{'proof':proof},format='json'))
        self.assertEqual(sorted(results),[201,409]);self.assertEqual(AdvisorAssignment.objects.count(),1)
    def test_simultaneous_changes_cannot_overwrite_a_new_revision(self):
        from .models import AdvisorAssignment,BusinessAudit
        row=AdvisorAssignment.objects.create(setter_id=21,setter_name='Setter',advisor_id=31,advisor_name='Advisor')
        results=self.race(lambda client:client.patch(f'/api/assurance/assignments/{row.pk}/',{'active':False,'revision':0},format='json'))
        self.assertEqual(sorted(results),[200,409])
        row.refresh_from_db();self.assertEqual(row.revision,1)
        self.assertEqual(BusinessAudit.objects.filter(action='assignment.updated').count(),1)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class AppointmentTests(TestCase):
    def setUp(self):
        from unittest.mock import patch
        external=patch('apps.assurance.backend.crm.appointments.external_check',return_value={'result':'internal_only','checked_at':'2026-10-21T10:00:00+00:00'})
        self.external_check=external.start();self.addCleanup(external.stop)
        from .models import AdvisorAssignment
        from django.utils import timezone
        from datetime import timedelta
        self.client=APIClient();self.authorize(21)
        self.contact=Contact.objects.create(name='Prospect',phone='raw',created_by_id=21)
        self.assignment=AdvisorAssignment.objects.create(setter_id=21,setter_name='SECRET SETTER',advisor_id=31,advisor_name='Advisor')
        self.token=self.client.post(f'/api/assurance/contacts/{self.contact.pk}/lease/',{},format='json').data['token']
        self.start=(timezone.now()+timedelta(days=10)).replace(hour=10,minute=0,second=0,microsecond=0)
        self.end=self.start+timedelta(minutes=30)
        self.url='/api/assurance/appointments/'
    def authorize(self,subject,role='setter'):
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(subject,'assurance',role))
    def proof(self,actor=21,setter=21,advisor=31):
        from packages.access.assignments import issue_booking
        return issue_booking(actor,setter,'SECRET SETTER',advisor,'Advisor')
    def payload(self,**changes):
        from uuid import uuid4
        return {'contact_id':self.contact.pk,'assignment_id':self.assignment.pk,'request_id':str(uuid4()),'token':self.token,'proof':self.proof(),
                'starts_at':self.start.isoformat(),'ends_at':self.end.isoformat(),'timezone':'UTC',**changes}
    def book(self,**changes):return self.client.post(self.url,self.payload(**changes),format='json')
    def reschedule_payload(self,**changes):
        from datetime import timedelta
        return {'revision':0,'proof':self.proof(),'starts_at':(self.start+timedelta(days=1)).isoformat(),
                'ends_at':(self.end+timedelta(days=1)).isoformat(),'timezone':'UTC','reason':'Internal SECRET SETTER note',**changes}
    def test_calendar_window_overlap_boundaries_and_visibility(self):
        from datetime import timedelta
        row=self.book();self.assertEqual(row.status_code,201,row.data)
        inside={'ends_after':(self.start+timedelta(minutes=5)).isoformat(),'starts_before':(self.end-timedelta(minutes=5)).isoformat()}
        self.assertEqual(self.client.get(self.url,inside).data['count'],1)
        self.assertEqual(self.client.get(self.url,{'ends_after':self.end.isoformat(),'starts_before':(self.end+timedelta(days=1)).isoformat()}).data['count'],0)
        self.assertEqual(self.client.get(self.url,{'ends_after':(self.start-timedelta(days=1)).isoformat(),'starts_before':self.start.isoformat()}).data['count'],0)
        self.authorize(22)
        self.assertEqual(self.client.get(self.url,inside).data['count'],0)

    def test_calendar_window_rejects_invalid_or_unbounded_ranges(self):
        from datetime import timedelta
        for params in ({'ends_after':self.start.isoformat()},
                       {'ends_after':self.start.replace(tzinfo=None).isoformat(),'starts_before':self.end.isoformat()},
                       {'ends_after':'invalid','starts_before':self.end.isoformat()},
                       {'ends_after':self.end.isoformat(),'starts_before':self.start.isoformat()},
                       {'ends_after':self.start.isoformat(),'starts_before':(self.start+timedelta(days=63)).isoformat()}):
            self.assertEqual(self.client.get(self.url,params).status_code,400,params)

    def test_booking_is_idempotent_and_keeps_one_responsible_setter(self):
        from .models import Appointment,BusinessAudit
        data=self.payload();first=self.client.post(self.url,data,format='json');retry=self.client.post(self.url,data,format='json')
        self.assertEqual(first.status_code,201,first.data);self.assertEqual(retry.status_code,200,retry.data)
        self.assertEqual(first.data['id'],retry.data['id']);self.assertEqual(Appointment.objects.count(),1)
        row=Appointment.objects.get();self.assertEqual(row.setter_id,21);self.assertEqual(row.booked_by_id,21)
        self.assertEqual(BusinessAudit.objects.filter(action='appointment.booked').count(),1)
    def test_booking_requires_owned_live_lease_active_assignment_and_correct_proof(self):
        from uuid import uuid4
        from packages.access.assignments import issue_assignment
        for changes in ({'token':str(uuid4())},{'proof':self.proof(actor=99)},{'proof':issue_assignment(21,21,'Setter',31,'Advisor')}):
            self.assertIn(self.book(**changes).status_code,(400,409))
        self.assignment.active=False;self.assignment.save()
        self.assertEqual(self.book().status_code,409)
    def test_other_setter_cannot_use_someone_elses_assignment(self):
        self.authorize(22)
        self.assertEqual(self.book(proof=self.proof(actor=22)).status_code,403)
    def test_invalid_intervals_naive_times_and_wrong_dst_offsets_are_rejected(self):
        for changes in ({'ends_at':self.start.isoformat()},{'starts_at':self.start.replace(tzinfo=None).isoformat()},
                        {'timezone':'Unknown/Zone'},{'starts_at':'2027-03-14T02:30:00-05:00','ends_at':'2027-03-14T04:00:00-04:00','timezone':'America/Toronto'}):
            self.assertEqual(self.book(**changes).status_code,400,changes)
    def test_overlap_is_blocked_but_adjacent_slot_is_allowed(self):
        self.assertEqual(self.book().status_code,201)
        self.assertEqual(self.book().status_code,409)
        from datetime import timedelta
        self.assertEqual(self.book(starts_at=self.end.isoformat(),ends_at=(self.end+timedelta(minutes=30)).isoformat()).status_code,201)
    def test_same_prospect_cannot_overlap_with_another_advisor(self):
        from .models import AdvisorAssignment
        self.book()
        second=AdvisorAssignment.objects.create(setter_id=21,setter_name='Setter',advisor_id=32,advisor_name='Other')
        self.assertEqual(self.book(assignment_id=second.pk,proof=self.proof(advisor=32)).status_code,409)
    def test_reschedule_preserves_history_author_and_rejects_stale_revision(self):
        from .models import Appointment,BusinessAudit
        row=self.book().data;url=self.url+str(row['id'])+'/reschedule/'
        response=self.client.post(url,self.reschedule_payload(),format='json')
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(self.client.post(url,self.reschedule_payload(),format='json').status_code,409)
        event=BusinessAudit.objects.get(action='appointment.rescheduled')
        self.assertEqual(event.metadata['before']['starts_at'],self.start.isoformat())
        self.assertEqual(event.metadata['reason'],'Internal SECRET SETTER note')
        appointment=Appointment.objects.get();self.assertEqual(appointment.setter_id,21);self.assertEqual(appointment.booked_by_id,21)
    def test_cancel_frees_slot_even_if_assignment_has_been_suspended(self):
        row=self.book().data;self.assignment.active=False;self.assignment.save()
        url=self.url+str(row['id'])+'/cancel/'
        self.assertEqual(self.client.post(url,{'revision':0,'reason':'Cancelled'},format='json').status_code,200)
        self.assignment.active=True;self.assignment.save()
        self.assertEqual(self.book().status_code,201)
        self.assertEqual(self.client.get(self.url+'?status=cancelled').data['count'],1)
    def test_advisor_only_sees_own_appointment_and_linked_prospect_without_setter_metadata(self):
        row=self.book().data;self.authorize(31,'advisor')
        response=self.client.get(self.url)
        self.assertEqual(response.data['count'],1)
        public=response.data['results'][0]
        self.assertEqual(set(public),{'id','contact','advisor','starts_at','ends_at','timezone','status','revision','calendar','qualification_snapshot'})
        self.assertNotIn('SECRET SETTER',str(public));self.assertNotIn('setter',str(public))
        self.assertEqual(self.client.get('/api/assurance/contacts/').data['count'],1)
        self.assertEqual(self.client.get('/api/assurance/dashboard/').data['contacts'],1)
        for suffix in ('cancel/','reschedule/'):
            self.assertEqual(self.client.post(self.url+str(row['id'])+'/'+suffix,{},format='json').status_code,403)
        self.assertEqual(self.client.get(self.url+str(row['id'])+'/history/').status_code,403)
        self.authorize(32,'advisor');self.assertEqual(self.client.get(self.url).data['count'],0)
        self.assertEqual(self.client.get('/api/assurance/contacts/').data['count'],0)
    def test_other_setter_cannot_view_or_change_an_appointment(self):
        row=self.book().data;self.authorize(22)
        url=self.url+str(row['id'])+'/'
        self.assertEqual(self.client.get(url).status_code,404)
        self.assertEqual(self.client.post(url+'cancel/',{'revision':0,'reason':'x'},format='json').status_code,404)
    def test_admin_history_contains_internal_author_but_not_other_appointments(self):
        row=self.book().data;self.authorize(10,'administrator')
        response=self.client.get(self.url+str(row['id'])+'/history/')
        self.assertEqual(response.data['count'],1)
        self.assertEqual(response.data['results'][0]['actor_id'],21)
        self.assertEqual(self.client.get(self.url).data['results'][0]['setter']['id'],21)
    def test_admin_booking_keeps_original_responsible_setter_separate_from_actor(self):
        from .models import Appointment
        self.client.delete(f'/api/assurance/contacts/{self.contact.pk}/lease/',{'token':self.token},format='json')
        self.authorize(10,'administrator')
        self.token=self.client.post(f'/api/assurance/contacts/{self.contact.pk}/lease/',{},format='json').data['token']
        response=self.book(proof=self.proof(actor=10))
        self.assertEqual(response.status_code,201,response.data)
        self.assertEqual(response.data['setter']['id'],21)
        row=Appointment.objects.get();self.assertEqual(row.booked_by_id,10);self.assertEqual(row.setter_id,21)
        changed=self.client.post(self.url+str(row.pk)+'/reschedule/',self.reschedule_payload(proof=self.proof(actor=10)),format='json')
        self.assertEqual(changed.status_code,200,changed.data)
        row.refresh_from_db();self.assertEqual(row.booked_by_id,10);self.assertEqual(row.setter_id,21)
    def test_audit_failure_rolls_back_booking(self):
        from unittest.mock import patch
        from .models import Appointment
        with patch('apps.assurance.backend.crm.appointments.BusinessAudit.objects.create',side_effect=RuntimeError('audit failure')):
            with self.assertRaises(RuntimeError):self.book()
        self.assertFalse(Appointment.objects.exists())

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class ConcurrentAppointmentTests(TransactionTestCase):
    def setUp(self):
        from unittest.mock import patch
        external=patch('apps.assurance.backend.crm.appointments.external_check',return_value={'result':'internal_only','checked_at':'2026-10-21T10:00:00+00:00'})
        self.external_check=external.start();self.addCleanup(external.stop)
        from django.db import connection
        from .models import AdvisorAssignment
        from django.utils import timezone
        from datetime import timedelta
        if connection.vendor!='postgresql':self.skipTest('Concurrence réelle vérifiée dans la CI PostgreSQL')
        self.contacts=[];self.assignments=[];self.tokens=[]
        self.start=(timezone.now()+timedelta(days=10)).replace(hour=10,minute=0,second=0,microsecond=0)
        self.end=self.start+timedelta(minutes=30)
        for subject in (21,22):
            contact=Contact.objects.create(name='Prospect',phone='raw',created_by_id=subject)
            assignment=AdvisorAssignment.objects.create(setter_id=subject,setter_name='Setter',advisor_id=31,advisor_name='Advisor')
            client=APIClient();client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(subject,'assurance','setter'))
            token=client.post(f'/api/assurance/contacts/{contact.pk}/lease/',{},format='json').data['token']
            self.contacts.append(contact);self.assignments.append(assignment);self.tokens.append(token)
    def payload(self,index):
        from uuid import uuid4
        from packages.access.assignments import issue_booking
        subject=21+index
        return {'contact_id':self.contacts[index].pk,'assignment_id':self.assignments[index].pk,'token':self.tokens[index],
            'request_id':str(uuid4()),'proof':issue_booking(subject,subject,'Setter',31,'Advisor'),
            'starts_at':self.start.isoformat(),'ends_at':self.end.isoformat(),'timezone':'UTC'}
    def race(self,operations):
        from django.db import connections
        from threading import Barrier
        from concurrent.futures import ThreadPoolExecutor
        barrier=Barrier(2)
        def execute(index):
            try:
                actor,operation=operations[index]
                client=APIClient();client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(actor,'assurance','setter'))
                barrier.wait(timeout=10)
                return operation(client).status_code
            finally:connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:return list(pool.map(execute,[0,1]))
    def test_different_setters_cannot_double_book_the_same_advisor(self):
        from .models import Appointment
        first,second=self.payload(0),self.payload(1)
        results=self.race([(21,lambda api:api.post('/api/assurance/appointments/',first,format='json')),
                           (22,lambda api:api.post('/api/assurance/appointments/',second,format='json'))])
        self.assertEqual(sorted(results),[201,409]);self.assertEqual(Appointment.objects.count(),1)
    def test_simultaneous_retry_of_one_request_creates_only_one_appointment(self):
        from .models import Appointment,BusinessAudit
        payload=self.payload(0)
        operation=lambda api:api.post('/api/assurance/appointments/',payload,format='json')
        results=self.race([(21,operation),(21,operation)])
        self.assertEqual(sorted(results),[200,201]);self.assertEqual(Appointment.objects.count(),1)
        self.assertEqual(BusinessAudit.objects.filter(action='appointment.booked').count(),1)
    def test_rescheduling_and_new_booking_for_same_slot_have_only_one_winner(self):
        from .models import Appointment
        from packages.access.assignments import issue_booking
        from datetime import timedelta
        client=APIClient();client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(21,'assurance','setter'))
        created=client.post('/api/assurance/appointments/',self.payload(0),format='json').data
        start=self.start+timedelta(hours=2);end=self.end+timedelta(hours=2)
        change={'starts_at':start.isoformat(),'ends_at':end.isoformat(),'timezone':'UTC','revision':0,'reason':'Report',
                'proof':issue_booking(21,21,'Setter',31,'Advisor')}
        other={**self.payload(1),'starts_at':start.isoformat(),'ends_at':end.isoformat()}
        results=self.race([(21,lambda api:api.post(f'/api/assurance/appointments/{created["id"]}/reschedule/',change,format='json')),
                           (22,lambda api:api.post('/api/assurance/appointments/',other,format='json'))])
        self.assertEqual(results.count(409),1);self.assertIn(results[0],(200,409));self.assertIn(results[1],(201,409))
        self.assertEqual(Appointment.objects.filter(status='scheduled',starts_at=start).count(),1)
    def test_same_contact_cannot_be_booked_with_two_advisors_simultaneously(self):
        from .models import AdvisorAssignment,Appointment
        from packages.access.assignments import issue_booking
        from uuid import uuid4
        first=self.payload(0)
        other_assignment=AdvisorAssignment.objects.create(setter_id=21,setter_name='Setter',advisor_id=32,advisor_name='Other')
        second={**first,'assignment_id':other_assignment.pk,'request_id':str(uuid4()),'proof':issue_booking(21,21,'Setter',32,'Other')}
        results=self.race([(21,lambda api:api.post('/api/assurance/appointments/',first,format='json')),
                           (21,lambda api:api.post('/api/assurance/appointments/',second,format='json'))])
        self.assertEqual(sorted(results),[201,409]);self.assertEqual(Appointment.objects.count(),1)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class ReportTests(TestCase):
    def setUp(self):
        from .models import Appointment,AdvisorAssignment
        from django.utils import timezone
        from datetime import timedelta
        from uuid import uuid4
        self.client=APIClient();self.auth(31,'advisor')
        contact=Contact.objects.create(name='Prospect',phone='raw',created_by_id=21)
        assignment=AdvisorAssignment.objects.create(setter_id=21,setter_name='SECRET',advisor_id=31,advisor_name='Advisor')
        self.appointment=Appointment.objects.create(request_id=uuid4(),fingerprint='x',contact=contact,assignment=assignment,setter_id=21,setter_name='SECRET',advisor_id=31,advisor_name='Advisor',booked_by_id=99,starts_at=timezone.now()-timedelta(hours=2),ends_at=timezone.now()-timedelta(hours=1),timezone='UTC')
        self.url=f'/api/assurance/appointments/{self.appointment.pk}/report/'
    def auth(self,pk,role):self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(pk,'assurance',role))
    def report(self,**values):return self.client.post(self.url,{'revision':0,'performed':False,'problem':False,'note':'No show',**values},format='json')
    def test_negative_report_creates_one_task_for_original_setter_and_retry_is_idempotent(self):
        from .models import FollowUpTask
        self.assertEqual(self.report().status_code,201)
        self.assertEqual(self.report().status_code,200)
        task=FollowUpTask.objects.get();self.assertEqual(task.setter_id,21)
        self.assertEqual(task.appointment_id,self.appointment.pk)
        self.assertNotIn('SECRET',str(self.client.get(self.url).data))
    def test_positive_report_has_no_task_unless_problem(self):
        from .models import FollowUpTask
        self.assertEqual(self.report(performed=True,problem=True).status_code,201)
        self.assertEqual(FollowUpTask.objects.count(),1)
    def test_yes_without_problem_has_no_task_and_cannot_be_overwritten(self):
        from .models import FollowUpTask
        self.assertEqual(self.report(performed=True,note='').status_code,201)
        self.assertFalse(FollowUpTask.objects.exists())
        self.assertEqual(self.report().status_code,409)
    def test_missing_report_is_pending_not_a_negative_report(self):
        from .models import FollowUpTask
        data=self.client.get('/api/assurance/dashboard/').data
        self.assertEqual(data['reports_due'],1);self.assertFalse(FollowUpTask.objects.exists())
        self.report();self.assertEqual(self.client.get('/api/assurance/dashboard/').data['reports_due'],0)
    def test_other_advisor_and_setter_cannot_submit(self):
        self.auth(32,'advisor');self.assertEqual(self.report().status_code,404)
        self.auth(21,'setter');self.assertEqual(self.report().status_code,403)
        self.auth(99,'administrator');self.assertEqual(self.report().status_code,403)
    def test_future_cancelled_stale_and_missing_reason_are_rejected(self):
        from django.utils import timezone
        from datetime import timedelta
        self.assertEqual(self.report(note='').status_code,400)
        self.assertEqual(self.report(revision=99).status_code,409)
        self.appointment.ends_at=timezone.now()+timedelta(hours=1);self.appointment.save()
        self.assertEqual(self.report().status_code,409)
        self.appointment.status='cancelled';self.appointment.save();self.assertEqual(self.report().status_code,409)
    def test_task_scope_alert_persistence_and_justified_close(self):
        from .models import FollowUpTask
        self.report();task=FollowUpTask.objects.get();url=f'/api/assurance/follow-ups/{task.pk}/'
        self.assertEqual(self.client.get('/api/assurance/follow-ups/').status_code,403)
        self.auth(22,'setter');self.assertEqual(self.client.get(url).status_code,404)
        self.auth(21,'setter')
        self.assertEqual(self.client.get('/api/assurance/dashboard/').data['open_follow_ups'],1)
        self.client.get(url)
        self.assertEqual(self.client.get('/api/assurance/dashboard/').data['open_follow_ups'],1)
        self.assertEqual(self.client.post(url+'close/',{'revision':0,'resolution':' '},format='json').status_code,400)
        self.assertEqual(self.client.post(url+'close/',{'revision':0,'resolution':'Prospect rappelé, rendez-vous à reprendre'},format='json').status_code,200)
        self.assertEqual(self.client.get('/api/assurance/dashboard/').data['open_follow_ups'],0)
        self.assertEqual(self.client.get('/api/assurance/follow-ups/?state=closed').data['count'],1)
        self.assertEqual(self.client.post(url+'close/',{'revision':0,'resolution':'Again'},format='json').status_code,404)
    def test_reported_appointment_cannot_be_cancelled_to_erase_outcome(self):
        self.report();self.auth(21,'setter')
        self.assertEqual(self.client.post(f'/api/assurance/appointments/{self.appointment.pk}/cancel/',{'revision':0,'reason':'Erase'},format='json').status_code,409)
    def test_audit_failure_rolls_back_report_and_task(self):
        from unittest.mock import patch
        from .models import AppointmentReport,FollowUpTask
        with patch('apps.assurance.backend.crm.reports.BusinessAudit.objects.create',side_effect=RuntimeError('audit')):
            with self.assertRaises(RuntimeError):self.report()
        self.assertFalse(AppointmentReport.objects.exists());self.assertFalse(FollowUpTask.objects.exists())

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class ConcurrentReportTests(TransactionTestCase):
    def test_simultaneous_reports_create_only_one_task(self):
        from django.db import connection,connections
        from .models import Appointment,AdvisorAssignment,AppointmentReport,FollowUpTask
        from django.utils import timezone
        from datetime import timedelta
        from uuid import uuid4
        from threading import Barrier
        from concurrent.futures import ThreadPoolExecutor
        if connection.vendor!='postgresql':self.skipTest('Concurrence vérifiée dans la CI PostgreSQL')
        contact=Contact.objects.create(name='Prospect',phone='raw',created_by_id=21)
        assignment=AdvisorAssignment.objects.create(setter_id=21,setter_name='Setter',advisor_id=31,advisor_name='Advisor')
        row=Appointment.objects.create(request_id=uuid4(),fingerprint='x',contact=contact,assignment=assignment,setter_id=21,setter_name='Setter',advisor_id=31,advisor_name='Advisor',booked_by_id=21,starts_at=timezone.now()-timedelta(hours=2),ends_at=timezone.now()-timedelta(hours=1),timezone='UTC')
        barrier=Barrier(2)
        def submit(_):
            try:
                client=APIClient();client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(31,'assurance','advisor'))
                barrier.wait(timeout=10)
                return client.post(f'/api/assurance/appointments/{row.pk}/report/',{'revision':0,'performed':False,'problem':True,'note':'Problème'},format='json').status_code
            finally:connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(submit,[0,1]))
        self.assertEqual(sorted(results),[200,201]);self.assertEqual(AppointmentReport.objects.count(),1);self.assertEqual(FollowUpTask.objects.count(),1)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class TaskPlanTests(TestCase):
    def setUp(self):
        fixture=ReportTests(methodName='test_negative_report_creates_one_task_for_original_setter_and_retry_is_idempotent')
        fixture.setUp();fixture.report()
        from .models import FollowUpTask
        self.client=fixture.client;self.task=FollowUpTask.objects.get()
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(99,'assurance','administrator'))
        self.url=f'/api/assurance/follow-ups/{self.task.pk}/'
    def payload(self,**changes):
        from django.utils import timezone
        from datetime import timedelta
        return {'revision':0,'priority':'urgent','due_at':(timezone.now()+timedelta(days=2)).isoformat(),'reason':'Client à rappeler rapidement',**changes}
    def test_plan_updates_revision_and_audits_before_after(self):
        data=self.payload();response=self.client.post(self.url+'plan/',data,format='json')
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.data['priority'],'urgent');self.assertEqual(response.data['revision'],1)
        history=self.client.get(self.url+'history/').data
        self.assertEqual(history['count'],1);self.assertEqual(history['results'][0]['metadata']['before']['priority'],'normal')
        self.assertEqual(self.client.post(self.url+'plan/',data,format='json').status_code,409)
    def test_invalid_priority_date_and_reason_leave_task_unchanged(self):
        for changes in ({'priority':'unknown'},{'due_at':'2020-01-01T00:00:00Z'},{'due_at':'2030-01-01T00:00:00'},{'reason':' '},{'setter_id':100}):
            self.assertEqual(self.client.post(self.url+'plan/',self.payload(**changes),format='json').status_code,400)
        self.task.refresh_from_db();self.assertEqual(self.task.revision,0)
    def test_setter_and_advisor_cannot_plan_or_read_audit(self):
        for pk,role in ((21,'setter'),(31,'advisor')):
            self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(pk,'assurance',role))
            self.assertEqual(self.client.post(self.url+'plan/',self.payload(),format='json').status_code,403)
            self.assertEqual(self.client.get(self.url+'history/').status_code,403)
    def test_closed_task_cannot_be_replanned(self):
        self.client.post(self.url+'close/',{'revision':0,'resolution':'Terminé'},format='json')
        self.assertEqual(self.client.post(self.url+'plan/?state=all',self.payload(revision=1),format='json').status_code,409)
    def test_overdue_indicator_does_not_close_or_hide_task(self):
        from django.utils import timezone
        from datetime import timedelta
        from .models import FollowUpTask
        FollowUpTask.objects.filter(pk=self.task.pk).update(due_at=timezone.now()-timedelta(hours=1))
        row=self.client.get(self.url).data
        self.assertTrue(row['overdue']);self.assertIsNone(row['closed_at'])
    def test_audit_failure_rolls_back_plan(self):
        from unittest.mock import patch
        with patch('apps.assurance.backend.crm.reports.BusinessAudit.objects.create',side_effect=RuntimeError('audit')):
            with self.assertRaises(RuntimeError):self.client.post(self.url+'plan/',self.payload(),format='json')
        self.task.refresh_from_db();self.assertEqual(self.task.revision,0);self.assertEqual(self.task.priority,'normal')

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class ConcurrentTaskPlanTests(TransactionTestCase):
    def test_simultaneous_plans_have_one_winner(self):
        from django.db import connection,connections
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        if connection.vendor!='postgresql':self.skipTest('Concurrence vérifiée dans la CI PostgreSQL')
        fixture=TaskPlanTests(methodName='test_plan_updates_revision_and_audits_before_after');fixture.setUp()
        payload=fixture.payload();url=fixture.url;barrier=Barrier(2)
        def submit(priority):
            try:
                api=APIClient();api.credentials(HTTP_AUTHORIZATION='Bearer '+issue(99,'assurance','administrator'))
                barrier.wait(timeout=10)
                return api.post(url+'plan/',{**payload,'priority':priority},format='json').status_code
            finally:connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(submit,['high','urgent']))
        self.assertEqual(sorted(results),[200,409])
        fixture.task.refresh_from_db();self.assertEqual(fixture.task.revision,1)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class IncidentTests(TestCase):
    def setUp(self):
        fixture=ReportTests(methodName='test_negative_report_creates_one_task_for_original_setter_and_retry_is_idempotent');fixture.setUp()
        self.client=fixture.client;self.appointment=fixture.appointment;self.report=fixture.report
        self.url=f'/api/assurance/appointments/{self.appointment.pk}/incidents/'
    def signal(self,**changes):
        from uuid import uuid4
        return self.client.post(self.url,{'request_id':str(uuid4()),'note':'Problème ultérieur',**changes},format='json')
    def test_requires_initial_report_and_does_not_modify_it(self):
        from .models import AppointmentReport,FollowUpTask
        self.assertEqual(self.signal().status_code,409)
        self.report(performed=True,note='Initial')
        self.assertEqual(self.signal().status_code,201)
        self.assertTrue(AppointmentReport.objects.get().performed)
        self.assertEqual(AppointmentReport.objects.get().note,'Initial')
        self.assertEqual(FollowUpTask.objects.get().setter_id,21)
    def test_retry_is_idempotent_and_changed_content_conflicts(self):
        from uuid import uuid4
        from .models import AppointmentIncident,FollowUpTask
        self.report();key=str(uuid4())
        self.assertEqual(self.signal(request_id=key).status_code,201)
        self.assertEqual(self.signal(request_id=key).status_code,200)
        self.assertEqual(self.signal(request_id=key,note='Different').status_code,409)
        self.assertEqual(AppointmentIncident.objects.count(),1);self.assertEqual(FollowUpTask.objects.count(),1)
    def test_closed_task_is_reopened_with_old_resolution_in_audit(self):
        from .models import FollowUpTask,BusinessAudit
        self.report();task=FollowUpTask.objects.get()
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(21,'assurance','setter'))
        self.client.post(f'/api/assurance/follow-ups/{task.pk}/close/',{'revision':0,'resolution':'Ancienne clôture'},format='json')
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(31,'assurance','advisor'))
        self.assertEqual(self.signal().status_code,201)
        task.refresh_from_db();self.assertIsNone(task.closed_at);self.assertEqual(task.revision,2)
        self.assertEqual(BusinessAudit.objects.get(action='appointment.follow_up_incident').metadata['before']['resolution'],'Ancienne clôture')
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(21,'assurance','setter'))
        self.assertEqual(self.client.post(f'/api/assurance/follow-ups/{task.pk}/close/',{'revision':1,'resolution':'Stale'},format='json').status_code,409)
        self.assertEqual(self.client.get('/api/assurance/dashboard/').data['open_follow_ups'],1)
    def test_other_roles_or_advisor_cannot_submit_or_discover_setter(self):
        self.report();self.signal()
        public=self.client.get(self.url).data['results'][0]
        self.assertEqual(set(public),{'id','note','created_at'})
        self.assertNotIn('SECRET',str(public))
        for pk,role,expected in ((32,'advisor',404),(21,'setter',403),(99,'administrator',403)):
            self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(pk,'assurance',role))
            self.assertEqual(self.signal().status_code,expected)
    def test_multiple_incidents_keep_one_task_and_are_visible_in_admin_history(self):
        from .models import FollowUpTask
        self.report();self.signal();self.signal(note='Second problème')
        self.assertEqual(FollowUpTask.objects.count(),1)
        self.assertEqual(self.client.get(self.url).data['count'],2)
        task=FollowUpTask.objects.get();self.assertEqual(task.revision,2)
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(99,'assurance','administrator'))
        history=self.client.get(f'/api/assurance/follow-ups/{task.pk}/history/').data
        self.assertEqual(sum(row['action']=='appointment.incident_reported' for row in history['results']),2)
    def test_audit_failure_rolls_back_incident_and_reopening(self):
        from unittest.mock import patch
        from .models import AppointmentIncident,FollowUpTask
        self.report()
        with patch('apps.assurance.backend.crm.reports.BusinessAudit.objects.create',side_effect=RuntimeError('audit')):
            with self.assertRaises(RuntimeError):self.signal()
        self.assertFalse(AppointmentIncident.objects.exists());self.assertEqual(FollowUpTask.objects.get().revision,0)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class ConcurrentIncidentTests(TransactionTestCase):
    def test_simultaneous_retries_reopen_only_once(self):
        from django.db import connection,connections
        from .models import FollowUpTask,AppointmentIncident
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        from uuid import uuid4
        if connection.vendor!='postgresql':self.skipTest('Concurrence vérifiée dans la CI PostgreSQL')
        fixture=IncidentTests(methodName='test_requires_initial_report_and_does_not_modify_it');fixture.setUp();fixture.report()
        task=FollowUpTask.objects.get()
        from django.utils import timezone
        FollowUpTask.objects.filter(pk=task.pk).update(closed_at=timezone.now(),resolution='Ancienne clôture')
        key=str(uuid4());barrier=Barrier(2)
        def signal(_):
            try:
                api=APIClient();api.credentials(HTTP_AUTHORIZATION='Bearer '+issue(31,'assurance','advisor'))
                barrier.wait(timeout=10)
                return api.post(fixture.url,{'request_id':key,'note':'Nouveau problème'},format='json').status_code
            finally:connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(signal,[0,1]))
        self.assertEqual(sorted(results),[200,201]);self.assertEqual(AppointmentIncident.objects.count(),1)
        task.refresh_from_db();self.assertIsNone(task.closed_at);self.assertEqual(task.revision,1)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class FollowUpPolicyTests(TestCase):
    def setUp(self):
        self.client=APIClient();self.auth(99,'administrator');self.url='/api/assurance/follow-up-policy/'
    def auth(self,pk,role):self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(pk,'assurance',role))
    def payload(self,**changes):return {'initial_hours':48,'reopen_hours':6,'default_priority':'high','revision':0,'reason':'SLA client',**changes}
    def test_defaults_update_history_and_stale_write(self):
        self.assertEqual(self.client.get(self.url).data['initial_hours'],24)
        data=self.payload();response=self.client.post(self.url,data,format='json')
        self.assertEqual(response.status_code,200,response.data);self.assertEqual(response.data['revision'],1)
        self.assertEqual(self.client.post(self.url,data,format='json').status_code,409)
        history=self.client.get(self.url+'history/').data
        self.assertEqual(history['count'],1);self.assertEqual(history['results'][0]['metadata']['before']['initial_hours'],24)
    def test_roles_and_invalid_configuration(self):
        for changes in ({'initial_hours':0},{'reopen_hours':721},{'default_priority':'bad'},{'reason':' '},{'unexpected':True}):
            self.assertEqual(self.client.post(self.url,self.payload(**changes),format='json').status_code,400)
        for pk,role in ((21,'setter'),(31,'advisor')):
            self.auth(pk,role)
            self.assertEqual(self.client.get(self.url).status_code,403)
            self.assertEqual(self.client.get(self.url+'history/').status_code,403)
            self.assertEqual(self.client.post(self.url,self.payload(),format='json').status_code,403)
    def test_new_tasks_use_policy_without_replanning_existing_tasks(self):
        from .models import FollowUpTask
        from django.utils import timezone
        fixture=ReportTests(methodName='test_negative_report_creates_one_task_for_original_setter_and_retry_is_idempotent');fixture.setUp();fixture.report()
        task=FollowUpTask.objects.get();old_due=task.due_at
        self.assertEqual(self.client.post(self.url,self.payload(),format='json').status_code,200)
        task.refresh_from_db();self.assertEqual(task.due_at,old_due);self.assertEqual(task.priority,'normal')
        from .models import Appointment
        from uuid import uuid4
        row=Appointment.objects.get();row.pk=None;row.request_id=uuid4();row.save()
        fixture.url=f'/api/assurance/appointments/{row.pk}/report/';fixture.report()
        latest=FollowUpTask.objects.order_by('-pk').first()
        self.assertEqual(latest.priority,'high');self.assertEqual(latest.policy_revision,1)
        self.assertAlmostEqual((latest.due_at-timezone.now()).total_seconds()/3600,48,places=2)
    def test_reopening_uses_new_delay_and_preserves_existing_priority(self):
        from .models import FollowUpTask
        from django.utils import timezone
        from uuid import uuid4
        fixture=IncidentTests(methodName='test_requires_initial_report_and_does_not_modify_it');fixture.setUp();fixture.report()
        task=FollowUpTask.objects.get();FollowUpTask.objects.filter(pk=task.pk).update(closed_at=timezone.now(),resolution='Done',priority='urgent')
        self.client.post(self.url,self.payload(),format='json')
        fixture.signal(request_id=str(uuid4()))
        task.refresh_from_db();self.assertIsNone(task.closed_at);self.assertEqual(task.priority,'urgent');self.assertEqual(task.policy_revision,1)
        self.assertAlmostEqual((task.due_at-timezone.now()).total_seconds()/3600,6,places=2)
    def test_audit_failure_rolls_back_rules(self):
        from unittest.mock import patch
        with patch('apps.assurance.backend.crm.policies.BusinessAudit.objects.create',side_effect=RuntimeError('audit')):
            with self.assertRaises(RuntimeError):self.client.post(self.url,self.payload(),format='json')
        self.assertEqual(self.client.get(self.url).data['revision'],0)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class ConcurrentPolicyTests(TransactionTestCase):
    def test_simultaneous_rule_updates_have_one_winner(self):
        from django.db import connection,connections
        from .models import FollowUpPolicy
        from threading import Barrier
        from concurrent.futures import ThreadPoolExecutor
        if connection.vendor!='postgresql':self.skipTest('Concurrence vérifiée dans la CI PostgreSQL')
        FollowUpPolicy.objects.create(pk=1);barrier=Barrier(2)
        def update(hours):
            try:
                api=APIClient();api.credentials(HTTP_AUTHORIZATION='Bearer '+issue(99,'assurance','administrator'))
                barrier.wait(timeout=10)
                return api.post('/api/assurance/follow-up-policy/',{'initial_hours':hours,'reopen_hours':6,'default_priority':'high','revision':0,'reason':'New rule'},format='json').status_code
            finally:connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(update,[12,48]))
        self.assertEqual(sorted(results),[200,409]);self.assertEqual(FollowUpPolicy.objects.get().revision,1)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class ContactNoteTests(TestCase):
    def setUp(self):
        self.client=APIClient();self.auth(21,'setter')
        self.contact=Contact.objects.create(name='Prospect',phone='raw',created_by_id=21,assigned_advisor_id=31)
        self.url=f'/api/assurance/contacts/{self.contact.pk}/notes/'
        self.token=self.client.post(f'/api/assurance/contacts/{self.contact.pk}/lease/',{},format='json').data['token']
    def auth(self,pk,role):self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(pk,'assurance',role))
    def add(self,**changes):
        from uuid import uuid4
        return self.client.post(self.url,{'request_id':str(uuid4()),'body':'Note interne SECRET SETTER','token':self.token,**changes},format='json')
    def test_internal_note_requires_current_lease_and_is_hidden_from_advisor(self):
        self.assertEqual(self.add().status_code,201)
        self.auth(22,'setter');self.assertEqual(self.add().status_code,409)
        self.assertEqual(self.client.get(self.url).data['count'],1)
        self.auth(31,'advisor');self.assertEqual(self.client.get(self.url).data['count'],0)
        self.assertNotIn('SECRET',str(self.client.get(self.url).data))
    def test_advisor_notes_are_self_and_admin_only(self):
        from .models import ContactNote
        self.auth(31,'advisor');response=self.add(body='Ma note personnelle')
        self.assertEqual(response.status_code,201,response.data)
        self.assertEqual(set(response.data),{'id','body','created_at','channel'})
        ContactNote.objects.create(contact=self.contact,request_id=__import__('uuid').uuid4(),author_id=32,channel='advisor',body='Other advisor')
        self.assertEqual(self.client.get(self.url).data['count'],1)
        self.auth(21,'setter');self.assertEqual(self.client.get(self.url).data['count'],0)
        self.auth(99,'administrator');self.assertEqual(self.client.get(self.url).data['count'],2)
        self.assertIn('author_id',self.client.get(self.url).data['results'][0])
    def test_idempotent_retry_and_conflicting_content(self):
        from uuid import uuid4
        from .models import ContactNote
        key=str(uuid4())
        self.assertEqual(self.add(request_id=key).status_code,201)
        self.assertEqual(self.add(request_id=key).status_code,200)
        self.assertEqual(self.add(request_id=key,body='Other').status_code,409)
        self.assertEqual(ContactNote.objects.count(),1)
    def test_body_bounds_unknown_visibility_and_append_only_routes(self):
        self.assertEqual(self.add(body=' ').status_code,400)
        self.assertEqual(self.add(body='x'*5001).status_code,400)
        self.assertEqual(self.add(channel='advisor').status_code,400)
        self.assertEqual(self.add(author_id=99).status_code,400)
        self.add();self.add(body='Correction ajoutée')
        self.assertEqual(self.client.get(self.url).data['count'],2)
        self.assertEqual(self.client.patch(self.url,{'body':'Overwrite'},format='json').status_code,405)
        self.assertEqual(self.client.delete(self.url).status_code,405)
    def test_unassigned_advisor_cannot_access_notes(self):
        self.auth(32,'advisor')
        self.assertEqual(self.client.get(self.url).status_code,404)
        self.assertEqual(self.add().status_code,404)
    def test_audit_failure_rolls_back_note(self):
        from unittest.mock import patch
        from .models import ContactNote
        with patch('apps.assurance.backend.crm.notes.BusinessAudit.objects.create',side_effect=RuntimeError('audit')):
            with self.assertRaises(RuntimeError):self.add()
        self.assertFalse(ContactNote.objects.exists())

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class ConcurrentContactNoteTests(TransactionTestCase):
    def test_simultaneous_advisor_retries_keep_one_note(self):
        from django.db import connection,connections
        from .models import ContactNote
        from uuid import uuid4
        from threading import Barrier
        from concurrent.futures import ThreadPoolExecutor
        if connection.vendor!='postgresql':self.skipTest('Concurrence vérifiée dans la CI PostgreSQL')
        row=Contact.objects.create(name='Prospect',phone='raw',created_by_id=21,assigned_advisor_id=31)
        key=str(uuid4());barrier=Barrier(2)
        def add(_):
            try:
                api=APIClient();api.credentials(HTTP_AUTHORIZATION='Bearer '+issue(31,'assurance','advisor'))
                barrier.wait(timeout=10)
                return api.post(f'/api/assurance/contacts/{row.pk}/notes/',{'request_id':key,'body':'Ma note'},format='json').status_code
            finally:connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(add,[0,1]))
        self.assertEqual(sorted(results),[200,201]);self.assertEqual(ContactNote.objects.count(),1)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class StatisticsTests(TestCase):
    def setUp(self):
        from .models import Appointment,AdvisorAssignment
        from django.utils import timezone
        from datetime import timedelta
        from uuid import uuid4
        self.client=APIClient();self.auth(99,'administrator')
        self.now=timezone.now();self.url='/api/assurance/statistics/'
        contact=Contact.objects.create(name='Prospect',phone='raw',created_by_id=21)
        assignment=AdvisorAssignment.objects.create(setter_id=21,setter_name='Setter secret',advisor_id=31,advisor_name='Advisor')
        self.base={'contact':contact,'assignment':assignment,'setter_id':21,'setter_name':'Setter secret','advisor_id':31,'advisor_name':'Advisor','booked_by_id':99,'timezone':'UTC','fingerprint':'x'}
        self.row=Appointment.objects.create(**self.base,request_id=uuid4(),starts_at=self.now-timedelta(hours=2),ends_at=self.now-timedelta(hours=1))
    def auth(self,pk,role):self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(pk,'assurance',role))
    def query(self):
        from datetime import timedelta
        return self.url+f'?from={(self.now-timedelta(days=1)).date()}&to={(self.now+timedelta(days=1)).date()}'
    def test_missing_report_is_not_counted_as_no_and_zero_rate_is_unknown(self):
        data=self.client.get(self.query()).data['totals']
        self.assertEqual(data['appointments'],1);self.assertEqual(data['reports_due'],1)
        self.assertEqual(data['not_performed'],0);self.assertIsNone(data['performed_rate'])
        self.assertEqual(data['open_tasks'],0)
    def test_multiple_incidents_do_not_multiply_appointments_reports_or_tasks(self):
        from .models import AppointmentReport,AppointmentIncident,FollowUpTask
        from uuid import uuid4
        from datetime import timedelta
        AppointmentReport.objects.create(appointment=self.row,performed=True,problem=True,author_id=31)
        for _ in range(3):AppointmentIncident.objects.create(appointment=self.row,request_id=uuid4(),author_id=31,note='Problem')
        FollowUpTask.objects.create(appointment=self.row,setter_id=21,reason='Problem',due_at=self.now-timedelta(hours=1))
        response=self.client.get(self.query());self.assertEqual(response.status_code,200,response.data)
        data=response.data['totals']
        self.assertEqual(data['appointments'],1);self.assertEqual(data['performed'],1);self.assertEqual(data['incidents'],3)
        self.assertEqual(data['open_tasks'],1);self.assertEqual(data['overdue_tasks'],1);self.assertEqual(data['performed_rate'],100)
        self.assertEqual(response.data['setters']['results'][0]['appointments'],1)
    def test_setter_scope_has_no_other_setters_or_identity_metadata(self):
        from .models import Appointment
        from uuid import uuid4
        Appointment.objects.create(**{**self.base,'setter_id':22,'setter_name':'Hidden setter'},request_id=uuid4(),starts_at=self.row.starts_at,ends_at=self.row.ends_at)
        self.auth(21,'setter');data=self.client.get(self.query()).data
        self.assertEqual(data['totals']['appointments'],1);self.assertNotIn('setters',data)
        self.assertNotIn('Hidden setter',str(data));self.assertNotIn('Setter secret',str(data))
        self.auth(23,'setter');self.assertEqual(self.client.get(self.query()).data['totals']['appointments'],0)
    def test_advisor_anonymous_and_other_workspace_are_denied(self):
        self.auth(31,'advisor');self.assertEqual(self.client.get(self.url).status_code,403)
        self.client.credentials();self.assertEqual(self.client.get(self.url).status_code,401)
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(99,'isolation','administrator'))
        self.assertEqual(self.client.get(self.url).status_code,401)
    def test_date_validation_and_excluded_appointments(self):
        for query in ('?from=bad','?from=2026-02-30','?from=2026-10-20&to=2026-10-01','?from=2020-01-01&to=2026-01-01','?from=9999-12-31&to=9999-12-31'):
            self.assertEqual(self.client.get(self.url+query).status_code,400)
        self.assertEqual(self.client.get(self.url+'?from=2000-01-01&to=2000-01-02').data['totals']['appointments'],0)
    def test_cancelled_is_not_missing_report_and_rates_use_reports_only(self):
        from .models import Appointment,AppointmentReport
        from uuid import uuid4
        self.row.status='cancelled';self.row.save()
        yes=Appointment.objects.create(**self.base,request_id=uuid4(),starts_at=self.row.starts_at,ends_at=self.row.ends_at)
        no=Appointment.objects.create(**self.base,request_id=uuid4(),starts_at=self.row.starts_at,ends_at=self.row.ends_at)
        AppointmentReport.objects.create(appointment=yes,performed=True,author_id=31)
        AppointmentReport.objects.create(appointment=no,performed=False,author_id=31)
        data=self.client.get(self.query()).data['totals']
        self.assertEqual(data['cancelled'],1);self.assertEqual(data['reports_due'],0)
        self.assertEqual(data['reported'],2);self.assertEqual(data['performed_rate'],50)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class StatisticsExportTests(TestCase):
    def setUp(self):
        fixture=StatisticsTests(methodName='test_missing_report_is_not_counted_as_no_and_zero_rate_is_unknown');fixture.setUp()
        self.client=fixture.client;self.row=fixture.row;self.base=fixture.base;self.url=fixture.query()
    def test_csv_bom_readable_headers_and_audit(self):
        import csv,io
        from .models import BusinessAudit
        response=self.client.get(self.url+'&export=csv')
        self.assertEqual(response.status_code,200)
        self.assertTrue(response.content.startswith(b'\xef\xbb\xbf'))
        rows=list(csv.reader(io.StringIO(response.content.decode('utf-8-sig')),delimiter=';'))
        self.assertEqual(rows[0][0],'Périmètre');self.assertEqual(rows[1][0],'Total autorisé')
        self.assertIn('Non calculable',rows[1]);self.assertEqual(rows[2][2],'Setter secret')
        self.assertEqual(response['Cache-Control'],'no-store')
        self.assertEqual(BusinessAudit.objects.filter(action='statistics.exported').count(),1)
    def test_xlsx_has_summary_filters_and_formula_names_are_text(self):
        from openpyxl import load_workbook
        from io import BytesIO
        self.row.setter_name='=HYPERLINK("bad")';self.row.save()
        response=self.client.get(self.url+'&export=xlsx')
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(response.content),data_only=False)
        self.assertEqual(book.sheetnames,['Synthèse','Par setter'])
        self.assertEqual(book['Par setter']['C2'].data_type,'s')
        self.assertEqual(book['Par setter']['C2'].value,'=HYPERLINK("bad")')
        self.assertEqual(book['Synthèse'].freeze_panes,'A2')
        self.assertTrue(book['Par setter'].auto_filter.ref)
    def test_csv_formula_name_is_neutralized(self):
        self.row.setter_name='  =1+1';self.row.save()
        response=self.client.get(self.url+'&export=csv')
        self.assertIn("'  =1+1",response.content.decode('utf-8-sig'))
    def test_setter_export_does_not_contain_other_setter_or_identity_columns_data(self):
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(21,'assurance','setter'))
        response=self.client.get(self.url+'&export=csv')
        self.assertEqual(response.status_code,200)
        self.assertNotIn('Setter secret',response.content.decode('utf-8-sig'))
        from openpyxl import load_workbook
        from io import BytesIO
        book=load_workbook(BytesIO(self.client.get(self.url+'&export=xlsx').content))
        self.assertEqual(book.sheetnames,['Synthèse'])
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(31,'assurance','advisor'))
        self.assertEqual(self.client.get(self.url+'&export=csv').status_code,403)
    def test_admin_export_contains_every_setter_beyond_page_one(self):
        from .models import Appointment
        from uuid import uuid4
        for i in range(26):Appointment.objects.create(**{**self.base,'setter_id':100+i,'setter_name':f'Setter {i}'},request_id=uuid4(),starts_at=self.row.starts_at,ends_at=self.row.ends_at)
        self.assertEqual(len(self.client.get(self.url).data['setters']['results']),25)
        import csv,io
        rows=list(csv.reader(io.StringIO(self.client.get(self.url+'&export=csv').content.decode('utf-8-sig')),delimiter=';'))
        self.assertEqual(len(rows),29)
    def test_invalid_export_and_dates_are_refused(self):
        self.assertEqual(self.client.get(self.url+'&export=pdf').status_code,400)
        self.assertEqual(self.client.get('/api/assurance/statistics/?from=bad&export=xlsx').status_code,400)

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class CustomFieldTests(TestCase):
    def setUp(self):
        self.client=APIClient();self.auth(99,'administrator')
        self.url='/api/assurance/contact-fields/'
        self.contact=Contact.objects.create(name='Prospect',phone='raw',created_by_id=21,assigned_advisor_id=31)
        self.values_url=f'/api/assurance/contacts/{self.contact.pk}/custom_fields/'
        self.fields=[{'key':'city','label':'Ville','type':'text','required':True,'active':True,'choices':[]},
            {'key':'consent','label':'Accord','type':'boolean','required':True,'active':True,'choices':[]},
            {'key':'date','label':'Date','type':'date','required':False,'active':True,'choices':[]},
            {'key':'kind','label':'Type','type':'choice','required':False,'active':True,'choices':['A','B']}]
        self.assertEqual(self.configure().status_code,200)
        self.auth(21,'setter');self.token=self.client.post(f'/api/assurance/contacts/{self.contact.pk}/lease/',{},format='json').data['token']
    def auth(self,pk,role):self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(pk,'assurance',role))
    def configure(self,**changes):return self.client.post(self.url,{'revision':0,'fields':self.fields,'reason':'Configuration',**changes},format='json')
    def save(self,**changes):return self.client.post(self.values_url,{'revision':0,'schema_revision':1,'token':self.token,'values':{'city':'Montréal','consent':False,'date':'2026-10-20','kind':'A'},**changes},format='json')
    def test_types_false_boolean_and_audit(self):
        from .models import BusinessAudit
        response=self.save();self.assertEqual(response.status_code,200,response.data)
        self.assertFalse(response.data['values']['consent']);self.assertEqual(response.data['revision'],1)
        audit=BusinessAudit.objects.get(action='contact.custom_fields_changed')
        self.assertEqual(audit.metadata['after']['city'],'Montréal')
    def test_required_unknown_and_invalid_types_are_rejected(self):
        for values in ({'city':'','consent':True},{'city':'X','consent':'yes'},{'city':'X','consent':False,'date':'2026-02-30'},
                       {'city':'X','consent':False,'kind':'C'},{'city':'X','consent':False,'secret':'x'}):
            self.assertEqual(self.save(values=values).status_code,400,values)
        self.contact.refresh_from_db();self.assertEqual(self.contact.custom_values,{})
    def test_requires_owned_lease_and_matching_contact_and_schema_versions(self):
        from uuid import uuid4
        self.assertEqual(self.save(token=str(uuid4())).status_code,409)
        self.assertEqual(self.save(revision=99).status_code,409)
        self.assertEqual(self.save(schema_revision=0).status_code,409)
        self.save();self.assertEqual(self.save().status_code,409)
    def test_advisor_has_no_values_schema_or_public_metadata(self):
        self.save();self.auth(31,'advisor')
        self.assertEqual(self.client.get(self.values_url).status_code,403)
        self.assertEqual(self.save().status_code,403)
        self.assertEqual(self.client.get(self.url).status_code,403)
        self.assertNotIn('custom_values',self.client.get(f'/api/assurance/contacts/{self.contact.pk}/').data)
    def test_schema_admin_only_immutable_types_and_archive_preserves_values(self):
        self.assertEqual(self.configure().status_code,403)
        self.save();self.auth(99,'administrator')
        self.assertEqual(self.configure(revision=1,fields=[]).status_code,400)
        changed=[{**row,'type':'date'} if row['key']=='city' else row for row in self.fields]
        self.assertEqual(self.configure(revision=1,fields=changed).status_code,400)
        archived=[{**row,'active':False} if row['key']=='city' else row for row in self.fields]
        self.assertEqual(self.configure(revision=1,fields=archived).status_code,200)
        self.auth(21,'setter')
        self.assertEqual(self.save(revision=1,schema_revision=2,values={'consent':True}).status_code,200)
        self.contact.refresh_from_db();self.assertEqual(self.contact.custom_values['city'],'Montréal')
    def test_configuration_duplicate_keys_and_stale_revision(self):
        self.auth(99,'administrator')
        self.assertEqual(self.configure().status_code,409)
        self.assertEqual(self.configure(revision=1,fields=self.fields+self.fields).status_code,400)
    def test_audit_failure_rolls_back_value_change(self):
        from unittest.mock import patch
        with patch('apps.assurance.backend.crm.custom_fields.BusinessAudit.objects.create',side_effect=RuntimeError('audit')):
            with self.assertRaises(RuntimeError):self.save()
        self.contact.refresh_from_db();self.assertEqual(self.contact.custom_values,{})

@override_settings(DEBUG=True)
@skipUnless(settings.ROOT_URLCONF=='apps.assurance.backend.urls','Service Assurance indépendant')
class ConcurrentCustomFieldTests(TransactionTestCase):
    def test_simultaneous_values_do_not_overwrite_current_revision(self):
        from django.db import connection,connections
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        if connection.vendor!='postgresql':self.skipTest('Concurrence vérifiée dans la CI PostgreSQL')
        fixture=CustomFieldTests(methodName='test_types_false_boolean_and_audit');fixture.setUp();barrier=Barrier(2)
        def save(city):
            try:
                api=APIClient();api.credentials(HTTP_AUTHORIZATION='Bearer '+issue(21,'assurance','setter'))
                barrier.wait(timeout=10)
                return api.post(fixture.values_url,{'revision':0,'schema_revision':1,'token':fixture.token,'values':{'city':city,'consent':False}},format='json').status_code
            finally:connections.close_all()
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(save,['Montréal','Québec']))
        self.assertEqual(sorted(results),[200,409]);fixture.contact.refresh_from_db();self.assertEqual(fixture.contact.revision,1)
