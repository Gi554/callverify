"""Raw Assurance imports. No phone validation, deduplication or provider calls."""
import csv
import hashlib
import io
import json
from pathlib import Path
from uuid import UUID
from zipfile import ZipFile,BadZipFile
from xml.etree.ElementTree import ParseError
from django.db import transaction,IntegrityError
from django.db.models import F
from .custom_fields import schema,validate_values,output as schema_output,Conflict
from .models import ContactFieldSchema
from rest_framework.decorators import api_view,parser_classes
from rest_framework.parsers import MultiPartParser,FormParser
from rest_framework.exceptions import PermissionDenied,ValidationError
from rest_framework.response import Response
from rest_framework.pagination import PageNumberPagination
from openpyxl import load_workbook
from .models import Contact,ContactImport,BusinessAudit

MAX_BYTES=5*1024*1024
MAX_ROWS=10000
MAX_COLUMNS=100
MAX_CELL=10000

def admin(request):
    if request.user.role!='administrator':
        raise PermissionDenied('Import réservé aux administrateurs Assurance')

def read_table(upload,has_header):
    if not upload or upload.size>MAX_BYTES:
        raise ValidationError('Fichier requis, limité à 5 Mo')
    raw=upload.read(MAX_BYTES+1)
    if len(raw)>MAX_BYTES:raise ValidationError('Fichier limité à 5 Mo')
    extension=Path(upload.name).suffix.lower()
    try:
        if extension=='.csv':
            try:text=raw.decode('utf-8-sig')
            except UnicodeDecodeError:text=raw.decode('cp1252')
            try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t')
            except csv.Error:dialect=csv.excel
            iterator=csv.reader(io.StringIO(text,newline=''),dialect)
            workbook=None
        elif extension=='.xlsx':
            with ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist())>30*1024*1024:
                    raise ValidationError('Classeur décompressé trop volumineux')
            workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
            iterator=workbook.worksheets[0].iter_rows(values_only=True)
        else:raise ValidationError('Formats acceptés : CSV ou XLSX')
        table=[]
        try:
            for index,values in enumerate(iterator,1):
                if index>MAX_ROWS+int(has_header):raise ValidationError('Maximum 10 000 lignes par fichier')
                if len(values)>MAX_COLUMNS:raise ValidationError('Maximum 100 colonnes')
                row=['' if value is None else str(value) for value in values]
                if any(len(cell)>MAX_CELL for cell in row):raise ValidationError('Une cellule dépasse 10 000 caractères')
                table.append((index,row))
        finally:
            if workbook:workbook.close()
    except ValidationError:raise
    except (ValueError,TypeError,csv.Error,BadZipFile,KeyError,UnicodeError,OverflowError,OSError,ParseError) as error:
        raise ValidationError('Fichier illisible ou non pris en charge') from error
    if not table:raise ValidationError('Fichier vide')
    width=max(len(row) for _,row in table)
    if not width:raise ValidationError('Fichier vide')
    headings=table.pop(0)[1] if has_header else []
    columns=[{'index':i,'label':headings[i] if i<len(headings) and headings[i] else f'Colonne {i+1}'} for i in range(width)]
    rows=[(index,row+['']*(width-len(row))) for index,row in table if any(value.strip() for value in row)]
    if not rows:raise ValidationError('Aucune ligne de données')
    return raw,columns,rows

def output(batch):
    return {'id':batch.pk,'filename':batch.filename,'created_at':batch.created_at,
            'total':batch.total,'imported':batch.imported,'rejected':batch.rejected,'report':batch.report}

@api_view(['GET','POST'])
@parser_classes([MultiPartParser,FormParser])
def imports(request):
    admin(request)
    if request.method=='GET':
        pager=PageNumberPagination();pager.page_size=20
        rows=pager.paginate_queryset(ContactImport.objects.order_by('-pk'),request)
        return pager.get_paginated_response([output(row) for row in rows])
    has_header=request.data.get('has_header','true')
    if has_header not in ('true','false'):raise ValidationError('Option en-tête invalide')
    upload=request.FILES.get('file')
    raw,columns,rows=read_table(upload,has_header=='true')
    mode=request.data.get('mode','preview')
    if mode=='preview':
        return Response({'columns':columns,'rows':[{'line':line,'values':values} for line,values in rows[:5]],
                         'total':len(rows),'custom_schema':schema_output(schema()),'sheet':'Première feuille du classeur' if upload.name.lower().endswith('.xlsx') else None})
    if mode!='confirm':raise ValidationError('Action inconnue')
    try:
        mapping=json.loads(request.data.get('mapping','{}'))
        request_id=UUID(request.data.get('request_id',''))
        custom_mapping=json.loads(request.data.get('custom_mapping','{}'))
    except (ValueError,TypeError,AttributeError):raise ValidationError('Correspondance ou identifiant d’import invalide')
    if not isinstance(mapping,dict) or set(mapping)-{'name','phone','email'}:
        raise ValidationError('Correspondance invalide')
    if 'phone' not in mapping:raise ValidationError('Sélectionnez la colonne téléphone')
    if any(type(value) is not int or value<0 or value>=len(columns) for value in mapping.values()):
        raise ValidationError('Colonne inconnue')
    if len(set(mapping.values()))!=len(mapping):raise ValidationError('Chaque champ doit utiliser une colonne distincte')
    if not isinstance(custom_mapping,dict) or len(custom_mapping)>30:raise ValidationError('Correspondance métier invalide')
    if any(type(value) is not int or value<0 or value>=len(columns) for value in custom_mapping.values()):raise ValidationError('Colonne métier inconnue')
    if len(set([*mapping.values(),*custom_mapping.values()]))!=len(mapping)+len(custom_mapping):raise ValidationError('Chaque champ doit utiliser une colonne distincte')
    revision=request.data.get('schema_revision','')
    if custom_mapping and (not isinstance(revision,str) or not revision.isdecimal() or len(revision)>10):raise ValidationError('Révision du schéma métier requise')
    canonical=[has_header,mapping] if not custom_mapping else [has_header,mapping,custom_mapping,int(revision)]
    digest=hashlib.sha256(raw+json.dumps(canonical,sort_keys=True).encode()).hexdigest()
    def replay(batch):
        if batch.created_by_id!=request.user.pk or batch.fingerprint!=digest:
            return Response({'detail':'Identifiant d’import déjà utilisé'},status=409)
        return Response(output(batch))
    existing=ContactImport.objects.filter(request_id=request_id).first()
    if existing:return replay(existing)
    rules=schema();active={item['key']:item for item in rules.fields if item['active']}
    if custom_mapping:
        if rules.revision!=int(revision):raise Conflict('Configuration métier modifiée : refaites l’aperçu')
        if set(custom_mapping)-set(active):raise ValidationError('Champ métier inconnu ou archivé')
    accepted=[];errors=[]
    for line,values in rows:
        fields={field:values[index] for field,index in mapping.items()}
        reason='Téléphone manquant' if not fields['phone'].strip() else ''
        if any(len(value)>(254 if field=='email' else 250) for field,value in fields.items()):
            reason='Champ trop long (nom/téléphone : 250 ; e-mail : 254 caractères)'
        custom={}
        if not reason and custom_mapping:
            try:
                for key,index in custom_mapping.items():
                    value=values[index];item=active[key]
                    if value=='':value=None
                    elif item['type']=='boolean':
                        word=value.strip().casefold()
                        if word not in ('true','false','oui','non','1','0'):raise ValidationError({key:'Booléen : oui/non, true/false ou 1/0 requis'})
                        value=word in ('true','oui','1')
                    elif item['type']=='date' and len(value)==19 and value.endswith(' 00:00:00'):value=value[:10]
                    custom[key]=value
                validate_values(custom,list(active.values()))
            except ValidationError as error:
                reason='Champs métier : '+str(error.detail)
        if reason:errors.append({'line':line,'reason':reason})
        else:accepted.append(Contact(name=fields.get('name',''),phone=fields['phone'],email=fields.get('email',''),
                                     created_by_id=request.user.pk,source_row=line,source_values=values,custom_values=custom))
    try:
        with transaction.atomic():
            if custom_mapping:
                ContactFieldSchema.objects.filter(pk=1).update(revision=F('revision'))
                if schema().revision!=int(revision):raise Conflict('Configuration métier modifiée : refaites l’aperçu')
            batch=ContactImport.objects.create(request_id=request_id,fingerprint=digest,filename=Path(upload.name).name[:250],
                created_by_id=request.user.pk,total=len(rows),imported=len(accepted),rejected=len(errors),report=errors,columns=columns,mapping=mapping if not custom_mapping else {**mapping,"custom_fields":custom_mapping,"schema_revision":int(revision)})
            for contact in accepted:contact.import_batch=batch
            Contact.objects.bulk_create(accepted,batch_size=200)
            BusinessAudit.objects.create(actor_id=request.user.pk,action='contacts.imported',object_id=batch.pk)
    except IntegrityError:
        existing=ContactImport.objects.filter(request_id=request_id).first()
        if existing:return replay(existing)
        raise
    return Response(output(batch),status=201)
