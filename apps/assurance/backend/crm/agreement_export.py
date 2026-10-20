"""Administrative Excel export, using exactly the registry's filters and ordering."""
from io import BytesIO
import re
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from rest_framework.decorators import api_view
from rest_framework.exceptions import ValidationError
from .agreements import administrator
from .agreement_registry import Filters, filtered, deadline_state

MAX_ROWS = 10000
STATES = {'draft': 'Brouillon', 'tracking': 'En suivi', 'archived': 'Archivé'}
DEADLINES = {'overdue': 'Date de fin dépassée', 'upcoming': 'Échéance proche',
             'undated': 'Date de fin manquante', 'later': 'Échéance ultérieure', 'not_tracked': 'Hors suivi'}
LABELS = {'search': 'Référence ou prospect', 'kind': 'Type', 'status': 'État',
          'starts_from': 'Début à partir du', 'starts_to': 'Début jusqu’au',
          'ends_from': 'Fin à partir du', 'ends_to': 'Fin jusqu’au',
          'deadline': 'Alerte', 'horizon_days': 'Horizon (jours)'}

def text(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '\uFFFD', value) if isinstance(value, str) else value

@api_view(['GET'])
def export(request):
    administrator(request)
    if 'page' in request.query_params:
        raise ValidationError('L’export porte sur tous les résultats filtrés, sans pagination.')
    serializer = Filters(data=request.query_params)
    serializer.is_valid(raise_exception=True)
    data = serializer.validated_data
    qs, today, days, _ = filtered(data)
    # A single bounded read avoids a count/read race and never silently truncates.
    rows = list(qs[:MAX_ROWS + 1])
    if len(rows) > MAX_ROWS:
        raise ValidationError(f'Plus de {MAX_ROWS} suivis : précisez les filtres avant d’exporter.')
    book = Workbook(); sheet = book.active; sheet.title = 'Mandats et contrats'
    sheet.append(['Référence', 'Prospect', 'Réf. fiche', 'Type', 'État', 'Début', 'Fin',
                  'Échéance', 'Jours avant fin', 'Pièces liées'])
    for row in rows:
        sheet.append([text(row.reference), text(row.contact.name), row.contact_id,
                      'Mandat' if row.kind == 'mandate' else 'Contrat', STATES[row.status],
                      row.starts_on, row.ends_on, DEADLINES[deadline_state(row, today, days)],
                      (row.ends_on - today).days if row.ends_on else None, row.document_count])
    sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
    for i, width in enumerate([28, 32, 14, 16, 16, 16, 16, 30, 20, 16], 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row[5:7]:cell.number_format = 'dd/mm/yyyy'
    meta = book.create_sheet('Périmètre')
    meta.append(['Paramètre', 'Valeur'])
    meta.append(['Date de référence (UTC)', today])
    meta['B2'].number_format = 'dd/mm/yyyy'
    meta.append(['Nombre de suivis', len(rows)])
    meta.append(['Portée', 'Tous les résultats filtrés, au-delà de la page affichée.'])
    meta.append(['Échéances', 'Jours calendaires UTC ; alertes réservées aux états En suivi.'])
    meta.append(['Limite', 'Suivi administratif, sans signature ni validation juridique.'])
    for key, label in LABELS.items():
        value = data.get(key) or 'Tous / non renseigné'
        if key == 'kind':value = {'mandate': 'Mandat', 'contract': 'Contrat'}.get(value, value)
        if key == 'status':value = STATES.get(value, value)
        if key == 'deadline':value = DEADLINES.get(value, value)
        meta.append([label, text(value)])
    meta.column_dimensions['A'].width = 30; meta.column_dimensions['B'].width = 80
    meta.freeze_panes = 'A2'
    for tab in book:
        for cell in tab[1]:
            cell.font = Font(bold=True, color='FFFFFF');cell.fill = PatternFill('solid', fgColor='53661C')
        for row in tab:
            for cell in row:
                # Explicit text cells preserve literal values and block spreadsheet formulas.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    stream = BytesIO();book.save(stream)
    response = HttpResponse(stream.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="mandats-contrats-{today}.xlsx"'
    response['Cache-Control'] = 'no-store';response['X-Content-Type-Options'] = 'nosniff'
    return response
