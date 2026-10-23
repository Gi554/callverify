import csv
from unittest.mock import patch
from io import BytesIO, StringIO
from openpyxl import load_workbook
from django.test import TestCase, override_settings
from rest_framework.test import APIClient
from packages.access.tickets import issue
from .models import Contact, ContactFieldSchema, BusinessAudit


@override_settings(DEBUG=True)
class ContactExportTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.url = '/api/assurance/contacts/export/'
        ContactFieldSchema.objects.create(
            pk=1,
            fields=[
                {'key': 'budget', 'label': 'Budget', 'type': 'text', 'advisor_visible': True},
                {'key': 'secret', 'label': 'Note interne', 'type': 'text'},
                {'key': 'old', 'label': 'Ancien', 'type': 'text', 'active': False},
            ],
        )
        self.alice = Contact.objects.create(
            name='Alice Martin',
            phone='+33 6 12 34 56 78',
            email='alice@example.test',
            created_by_id=21,
            assigned_advisor_id=31,
            custom_values={'budget': '1200', 'secret': 'interne', 'old': 'x'},
        )
        self.bob = Contact.objects.create(
            name='=HYPERLINK("http://x")',
            phone='0612',
            status='callback',
            created_by_id=21,
        )

    def auth(self, pk, role):
        self.client.credentials(HTTP_AUTHORIZATION='Bearer ' + issue(pk, 'assurance', role))

    def csv_rows(self, query):
        response = self.client.get(self.url, {'file': 'csv', **query})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Cache-Control'], 'no-store')
        self.assertIn('prospects-assurance-', response['Content-Disposition'])
        text = response.content.decode('utf-8')
        self.assertTrue(text.startswith('﻿'))
        return list(csv.reader(StringIO(text[1:]), delimiter=';'))

    def test_setter_csv_follows_list_filters_and_neutralises_formulas(self):
        self.auth(21, 'setter')
        rows = self.csv_rows({})
        self.assertEqual(
            rows[0],
            ['Réf. fiche', 'Nom', 'Téléphone', 'E-mail', 'Étape', 'Mis à jour le', 'Budget']
            + ['Note interne'],
        )
        self.assertEqual([row[0] for row in rows[1:]], [str(self.bob.pk), str(self.alice.pk)])
        self.assertEqual(rows[1][1], '\'=HYPERLINK("http://x")')
        self.assertEqual(rows[2][2], '+33 6 12 34 56 78')
        self.assertEqual(rows[2][6:], ['1200', 'interne'])
        self.assertEqual(len(self.csv_rows({'search': 'alice'})), 2)
        self.assertEqual(len(self.csv_rows({'status': 'callback'})), 2)
        audit = BusinessAudit.objects.filter(action='contacts.exported').last()
        self.assertEqual(audit.actor_id, 21)
        self.assertEqual(audit.metadata['status'], 'callback')

    def test_administrator_excel_keeps_formulas_as_text(self):
        self.auth(99, 'administrator')
        response = self.client.get(self.url, {'file': 'xlsx'})
        self.assertEqual(response.status_code, 200)
        self.assertIn('.xlsx', response['Content-Disposition'])
        sheet = load_workbook(BytesIO(response.content)).worksheets[0]
        self.assertEqual(sheet.title, 'Prospects')
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet['B2'].value, '=HYPERLINK("http://x")')
        self.assertEqual(sheet['B2'].data_type, 's')

    def test_advisor_gets_only_own_prospects_and_visible_fields(self):
        self.auth(31, 'advisor')
        rows = self.csv_rows({})
        self.assertEqual(rows[0][-1], 'Budget')
        self.assertNotIn('Note interne', rows[0])
        self.assertEqual([row[1] for row in rows[1:]], ['Alice Martin'])
        self.auth(32, 'advisor')
        self.assertEqual(len(self.csv_rows({})), 1)

    def test_invalid_requests_are_refused(self):
        self.auth(21, 'setter')
        self.assertEqual(self.client.get(self.url).status_code, 400)
        self.assertEqual(self.client.get(self.url, {'file': 'pdf'}).status_code, 400)
        self.assertEqual(self.client.get(self.url, {'file': 'csv', 'page': 2}).status_code, 400)
        self.client.credentials()
        self.assertIn(self.client.get(self.url, {'file': 'csv'}).status_code, (401, 403))

    def test_volume_limits_never_truncate(self):
        self.auth(21, 'setter')
        with patch.dict('packages.exports.MAX_ROWS', {'csv': 2, 'xlsx': 1}):
            self.assertEqual(len(self.csv_rows({})), 3)
            response = self.client.get(self.url, {'file': 'xlsx'})
            self.assertEqual(response.status_code, 400)
            self.assertIn('choisissez CSV', str(response.data))
