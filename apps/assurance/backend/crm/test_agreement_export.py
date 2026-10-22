from datetime import date, datetime, timezone
from io import BytesIO
from unittest.mock import patch
from openpyxl import load_workbook
from django.test import TestCase, override_settings
from . import test_agreement_registry as registry_tests


@override_settings(DEBUG=True)
class AgreementExportTests(TestCase):
    auth = registry_tests.RegistryTests.auth
    create = registry_tests.RegistryTests.create

    def setUp(self):
        registry_tests.RegistryTests.setUp(self)
        self.export_url = self.url + 'export/'

    def workbook(self, query=None):
        response = self.client.get(self.export_url, query or {})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Cache-Control'], 'no-store')
        self.assertEqual(response['X-Content-Type-Options'], 'nosniff')
        self.assertIn('.xlsx', response['Content-Disposition'])
        return load_workbook(BytesIO(response.content))

    def test_all_pages_and_private_data_excluded(self):
        for n in range(30):
            self.create(f'EXPORT-{n}')
        book = self.workbook()
        sheet = book.worksheets[0]
        self.assertEqual(sheet.max_row, 33)
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:J33')
        content = str(list(sheet.values)) + str(list(book.worksheets[1].values))
        for secret in ('PRIVATE', 'fingerprint', 'author_id'):
            self.assertNotIn(secret, content)
        from .models import BusinessAudit

        self.assertFalse(BusinessAudit.objects.exists())

    def test_filters_and_order_match_registry(self):
        for query in (
            {'kind': 'mandate'},
            {'search': 'alice', 'status': 'archived'},
            {'starts_from': '2026-10-01', 'ends_to': '2026-12-31'},
            {'search': 'missing'},
        ):
            expected = self.client.get(self.url, query).data['results']
            book = self.workbook(query)
            self.assertEqual(
                [row[0] for row in list(book.worksheets[0].values)[1:]],
                [row['reference'] for row in expected],
            )
        book = self.workbook({'kind': 'mandate'})
        self.assertEqual(book.worksheets[0]['F2'].value.date(), date(2026, 10, 1))
        self.assertEqual(book.worksheets[0]['F2'].number_format, 'dd/mm/yyyy')
        self.assertIn(('Type', 'Mandat'), list(book.worksheets[1].values))

    @patch(
        'apps.assurance.backend.crm.agreement_registry.timezone.now',
        return_value=datetime(2026, 10, 20, tzinfo=timezone.utc),
    )
    def test_deadline_and_horizon_filters(self, now):
        self.create('LATE', status='tracking', ends_on=date(2026, 10, 19))
        self.create('SOON', status='tracking', ends_on=date(2026, 10, 27))
        self.create('LATER', status='tracking', ends_on=date(2026, 10, 28))
        book = self.workbook({'deadline': 'upcoming', 'horizon_days': '7'})
        rows = list(book.worksheets[0].values)[1:]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], 'SOON')
        self.assertEqual(rows[0][7:9], ('Échéance proche', 7))

    def test_literal_formula_strings(self):
        self.first.reference = '=HYPERLINK("https://example.com")'
        self.first.save()
        self.contact.name = '\t=1+1'
        self.contact.save()
        sheet = self.workbook({'kind': 'mandate'}).worksheets[0]
        self.assertEqual(sheet['A2'].data_type, 's')
        self.assertEqual(sheet['A2'].value, self.first.reference)
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertEqual(sheet['B2'].value, self.contact.name)

    def test_bound_refuses_instead_of_truncating(self):
        with patch('apps.assurance.backend.crm.agreement_export.MAX_ROWS', 1):
            self.assertEqual(self.client.get(self.export_url).status_code, 400)
        with patch('apps.assurance.backend.crm.agreement_export.MAX_ROWS', 2):
            self.assertEqual(self.workbook().worksheets[0].max_row, 3)

    def test_xml_controls_are_replaced_without_export_failure(self):
        self.contact.name = 'Nom\x01test'
        self.contact.save()
        self.assertEqual(
            self.workbook({'kind': 'mandate'}).worksheets[0]['B2'].value, 'Nom\ufffdtest'
        )

    def test_permissions_and_strict_parameters(self):
        for query in (
            {'page': 1},
            {'extra': 'secret'},
            {'status': 'signed'},
            {'ends_from': 'bad'},
            {'starts_from': '2026-12-01', 'starts_to': '2026-10-01'},
        ):
            self.assertEqual(self.client.get(self.export_url, query).status_code, 400)
        self.assertEqual(
            self.client.get(self.export_url + '?kind=mandate&kind=contract').status_code, 400
        )
        self.assertEqual(self.client.post(self.export_url, {}, format='json').status_code, 405)
        for role in ('setter', 'advisor'):
            self.auth(21, role)
            self.assertEqual(self.client.get(self.export_url).status_code, 403)
        self.auth(99, 'administrator', 'isolation')
        self.assertEqual(self.client.get(self.export_url).status_code, 401)
        self.client.credentials()
        self.assertEqual(self.client.get(self.export_url).status_code, 401)
