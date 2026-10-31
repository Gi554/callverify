"""Readable, scoped statistics exports; never include contact details or free notes."""

import csv
import io
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

LABELS = {
    'appointments': 'Rendez-vous',
    'cancelled': 'Annulés',
    'upcoming': 'À venir',
    'reported': 'Comptes rendus',
    'performed': 'Effectués (Oui)',
    'not_performed': 'Non effectués (tous motifs)',
    'absent': 'Absences déclarées',
    'other_not_performed': 'Non effectués (autre motif)',
    'reports_due': 'Comptes rendus manquants',
    'problem_reports': 'Comptes rendus avec problème',
    'incidents': 'Signalements ultérieurs',
    'open_tasks': 'Tâches ouvertes',
    'overdue_tasks': 'Tâches en retard',
    'closed_tasks': 'Tâches clôturées',
    'performed_rate': 'Taux effectué / comptes rendus (%)',
}


def safe(value):
    if not isinstance(value, str):
        return value
    # Spreadsheet programs may treat leading whitespace before = as a formula.
    return (
        "'" + value
        if value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n'))
        else value
    )


def response(result, groups, kind):
    meta = [str(result['from']), str(result['to']), result['as_of'].isoformat()]
    headers = [
        'Périmètre',
        'Setter ID',
        'Setter',
        *LABELS.values(),
        'Du (Québec)',
        'Au inclus (Québec)',
        'Calculé le (UTC)',
        'Définition',
    ]

    def values(label, row):
        return [
            label,
            row.get('setter_id', ''),
            row.get('setter_name', ''),
            *[row[key] if row[key] is not None else 'Non calculable' for key in LABELS],
            *meta,
            result['definition'],
        ]

    rows = [values('Total autorisé', result['totals']), *[values('Setter', row) for row in groups]]
    filename = f"statistiques-assurance-{result['from']}-{result['to']}.{kind}"
    if kind == 'csv':
        stream = io.StringIO(newline='')
        stream.write('\ufeff')
        writer = csv.writer(stream, delimiter=';')
        writer.writerow(headers)
        writer.writerows([[safe(value) for value in row] for row in rows])
        content = stream.getvalue().encode('utf-8')
        content_type = 'text/csv; charset=utf-8'
    else:
        book = Workbook()
        summary = book.active
        summary.title = 'Synthèse'
        summary.append(['Statistiques Assurance', 'Valeur'])
        for label, value in [
            ('Du (Québec)', meta[0]),
            ('Au inclus (Québec)', meta[1]),
            ('Calculé le (UTC)', meta[2]),
            ('Définition', result['definition']),
        ]:
            summary.append([label, value])
        for key, label in LABELS.items():
            summary.append(
                [
                    label,
                    result['totals'][key]
                    if result['totals'][key] is not None
                    else 'Non calculable',
                ]
            )
        summary.column_dimensions['A'].width = 44
        summary.column_dimensions['B'].width = 75
        summary['B5'].alignment = Alignment(wrap_text=True, vertical='top')
        summary.row_dimensions[5].height = 70
        if groups:
            sheet = book.create_sheet('Par setter')
            sheet.append(headers)
            for row in rows[1:]:
                sheet.append(row)
            sheet.auto_filter.ref = sheet.dimensions
            for i in range(1, len(headers) + 1):
                sheet.column_dimensions[get_column_letter(i)].width = 26
        for sheet in book:
            sheet.freeze_panes = 'A2'
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='53661C')
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'  # Force text, even malicious names beginning '='.
        stream = io.BytesIO()
        book.save(stream)
        content = stream.getvalue()
        content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    reply = HttpResponse(content, content_type=content_type)
    reply['Content-Disposition'] = f'attachment; filename="{filename}"'
    reply['Cache-Control'] = 'no-store'
    reply['X-Content-Type-Options'] = 'nosniff'
    return reply
