"""Cross-module regression journey using public APIs and disposable test data only."""

import json
from datetime import datetime, timezone
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient
from packages.access.tickets import issue
from .models import Contact, ContactImport, ContactDocument, Agreement, BusinessAudit


@override_settings(DEBUG=True)
class AgreementJourneyTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.auth(99, 'administrator')
        clock = patch(
            'apps.assurance.backend.crm.agreement_registry.timezone.now',
            return_value=datetime(2026, 10, 20, tzinfo=timezone.utc),
        )
        clock.start()
        self.addCleanup(clock.stop)
        network = patch(
            'requests.sessions.Session.request',
            side_effect=AssertionError('No provider call allowed in this journey'),
        )
        self.network = network.start()
        self.addCleanup(network.stop)

    def auth(self, pk, role, space='assurance'):
        self.client.credentials(HTTP_AUTHORIZATION='Bearer ' + issue(pk, space, role))

    def ok(self, response, status=200):
        self.assertEqual(
            response.status_code, status, getattr(response, 'data', response.content[:200])
        )
        return response.data

    def journey(self, extension):
        table = [
            ['Nom', 'Téléphone', 'Email'],
            ['Recette Alpha', 'test-number', 'private@example.invalid'],
            ['Recette Rejet', '', ''],
            ['Recette Doublon', 'test-number', ''],
        ]
        if extension == 'csv':
            raw = '\n'.join(';'.join(row) for row in table).encode('utf-8')
        else:
            book = Workbook()
            for row in table:
                book.active.append(row)
            stream = BytesIO()
            book.save(stream)
            raw = stream.getvalue()

        def upload():
            return SimpleUploadedFile('recette.' + extension, raw)

        url = '/api/assurance/imports/'
        preview = self.ok(
            self.client.post(
                url, {'file': upload(), 'has_header': 'true', 'mode': 'preview'}, format='multipart'
            )
        )
        self.assertEqual(preview['total'], 3)
        self.assertFalse(Contact.objects.exists())
        key = str(uuid4())

        def confirm():
            return self.client.post(
                url,
                {
                    'file': upload(),
                    'has_header': 'true',
                    'mode': 'confirm',
                    'request_id': key,
                    'mapping': json.dumps({'name': 0, 'phone': 1, 'email': 2}),
                },
                format='multipart',
            )

        batch = self.ok(confirm(), 201)
        self.assertEqual((batch['imported'], batch['rejected']), (2, 1))
        self.assertEqual(batch['report'][0]['line'], 3)
        self.ok(confirm())
        self.assertEqual(ContactImport.objects.count(), 1)
        self.assertEqual(
            Contact.objects.filter(phone='test-number').count(), 2
        )  # Raw import deliberately preserves duplicates.
        contact = self.ok(self.client.get('/api/assurance/contacts/?search=Recette%20Alpha'))[
            'results'
        ][0]
        base = f'/api/assurance/contacts/{contact["id"]}/'
        token = self.ok(self.client.post(base + 'lease/', {}, format='json'), 201)['token']
        edited = self.ok(
            self.client.patch(
                base,
                {'token': token, 'revision': contact['revision'], 'name': 'Recette Alpha corrigée'},
                format='json',
            )
        )
        self.assertEqual(edited['revision'], 1)
        body = {
            'token': token,
            'contact_revision': 1,
            'request_id': str(uuid4()),
            'kind': 'contract',
            'reference': 'RECETTE-001',
            'status': 'tracking',
            'starts_on': '2026-10-01',
            'ends_on': '2026-10-27',
            'reason': 'Recette automatisée',
        }
        created = self.ok(self.client.post(base + 'agreements/', body, format='json'), 201)
        aid = created['agreement']['id']
        self.assertEqual(created['contact_revision'], 2)
        self.ok(self.client.post(base + 'agreements/', body, format='json'))
        self.assertEqual(Agreement.objects.count(), 1)
        file_id = str(uuid4())
        binary = b'%PDF-1.4\n% disposable regression fixture\n%%EOF'

        def upload_doc():
            return self.client.post(
                base + 'documents/',
                {
                    'token': token,
                    'request_id': file_id,
                    'file': SimpleUploadedFile(
                        'piece-privee.pdf', binary, content_type='application/pdf'
                    ),
                },
                format='multipart',
            )

        document = self.ok(upload_doc(), 201)
        self.ok(upload_doc())
        self.assertEqual(ContactDocument.objects.count(), 1)
        link = {
            'token': token,
            'contact_revision': 2,
            'revision': 0,
            'document_id': document['id'],
            'enabled': True,
            'reason': 'Liaison de recette',
        }
        self.ok(self.client.post(base + f'agreements/{aid}/documents/', link, format='json'))
        self.assertEqual(
            self.client.get(base + f'documents/{document["id"]}/download/').content, binary
        )
        filters = {'search': 'RECETTE-001', 'deadline': 'upcoming', 'horizon_days': '7'}
        registry = self.ok(self.client.get('/api/assurance/agreement-registry/', filters))
        self.assertEqual(registry['count'], 1)
        self.assertEqual(registry['results'][0]['document_count'], 1)
        self.assertEqual(registry['statistics']['by_deadline']['upcoming'], 1)
        exported = self.client.get('/api/assurance/agreement-registry/export/', filters)
        self.assertEqual(exported.status_code, 200)
        sheet = load_workbook(BytesIO(exported.content)).worksheets[0]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['B2'].value, 'Recette Alpha corrigée')
        self.assertEqual(sheet['I2'].value, 7)
        self.assertEqual(sheet['J2'].value, 1)
        self.assertNotIn('private@example.invalid', str(list(sheet.values)))
        change = {key: value for key, value in body.items() if key not in ('kind', 'request_id')}
        change.update(
            contact_revision=3, revision=1, status='archived', reason='Archivage de recette'
        )
        self.ok(self.client.post(base + f'agreements/{aid}/edit/', change, format='json'))
        self.assertEqual(
            self.ok(self.client.get('/api/assurance/agreement-registry/', filters))['count'], 0
        )
        archive = self.ok(
            self.client.get('/api/assurance/agreement-registry/', {'status': 'archived'})
        )
        self.assertEqual(archive['statistics']['by_deadline']['not_tracked'], 1)
        history = self.ok(self.client.get(base + f'agreements/{aid}/history/'))
        self.assertEqual(history['count'], 3)
        self.assertEqual(history['results'][0]['after']['status'], 'archived')
        self.assertEqual(history['results'][0]['after']['document_ids'], [document['id']])
        stale = self.client.post(base + f'agreements/{aid}/documents/', link, format='json')
        self.assertEqual(stale.status_code, 409)
        # Adding a document or agreement never implicitly shares it with an advisor.
        Contact.objects.filter(pk=contact['id']).update(assigned_advisor_id=31)
        self.auth(31, 'advisor')
        self.assertEqual(self.ok(self.client.get(base + 'documents/'))['count'], 0)
        self.assertEqual(
            self.client.get(base + f'documents/{document["id"]}/download/').status_code, 404
        )
        self.assertEqual(self.client.get('/api/assurance/agreement-registry/').status_code, 403)
        self.assertEqual(
            self.client.get('/api/assurance/agreement-registry/export/').status_code, 403
        )
        self.auth(99, 'administrator')
        self.assertTrue(BusinessAudit.objects.filter(action='contacts.imported').exists())
        self.assertEqual(
            self.client.delete(base + 'lease/', {'token': token}, format='json').status_code, 204
        )
        self.assertFalse(self.ok(self.client.get(base + 'lease/'))['locked'])
        self.network.assert_not_called()

    def test_csv_import_to_private_documents_statistics_and_excel(self):
        self.journey('csv')

    def test_xlsx_import_to_private_documents_statistics_and_excel(self):
        self.journey('xlsx')
