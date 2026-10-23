import csv
from io import StringIO, BytesIO
from unittest.mock import patch
from uuid import uuid4
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from . import tests as fixtures
from .models import Contact, AppointmentReport, ContactCallOutcome, BusinessAudit


@override_settings(DEBUG=True)
class ClientFolderExportTests(TestCase):
    def setUp(self):
        self.fixture = fixtures.ReportTests()
        self.fixture.setUp()
        self.fixture.auth(99, 'administrator')
        self.client = self.fixture.client
        self.row = self.fixture.appointment
        self.url = '/api/assurance/client-folders/31/export/'
        self.literal = Contact.objects.create(
            name='  =1+1', phone='+15145550100', created_by_id=21, assigned_advisor_id=31
        )
        Contact.objects.create(
            name='Foreign secret', phone='private', created_by_id=21, assigned_advisor_id=32
        )

    def csv(self, **query):
        result = self.client.get(self.url, {'file': 'csv', **query})
        self.assertEqual(result.status_code, 200)
        return result, list(
            csv.DictReader(StringIO(result.content.decode('utf-8-sig')), delimiter=';')
        )

    def test_exports_all_filtered_rows_with_literal_formulas_and_scoped_audit(self):
        result, rows = self.csv(section='prospects')
        self.assertEqual(len(rows), 2)
        self.assertNotIn('Foreign secret', str(rows))
        self.assertEqual(rows[0]['Prospect'], "'  =1+1")
        self.assertEqual(rows[0]['Téléphone fourni'], '+15145550100')
        self.assertTrue(
            result['Content-Disposition'].startswith(
                'attachment; filename="dossier-assurance-31-prospects-'
            )
        )
        self.assertEqual(result['Cache-Control'], 'no-store')
        _, filtered = self.csv(section='prospects', search='Prospect')
        self.assertEqual(len(filtered), 1)
        audit = BusinessAudit.objects.filter(action='client_folder.exported').latest('pk')
        self.assertEqual(audit.actor_id, 99)
        self.assertEqual(
            audit.metadata, {'client_id': 31, 'section': 'prospects', 'file': 'csv', 'rows': 1}
        )
        self.assertNotIn('search', audit.metadata)

    def test_excel_literals_and_indicators_keep_the_same_attendance_rules(self):
        result = self.client.get(self.url, {'file': 'xlsx', 'section': 'prospects'})
        self.assertEqual(result.status_code, 200)
        sheet = load_workbook(BytesIO(result.content), data_only=False).active
        self.assertEqual(sheet.cell(2, 5).value, '  =1+1')
        self.assertEqual(sheet.cell(2, 5).data_type, 's')
        AppointmentReport.objects.create(
            appointment=self.row, performed=True, sale_concluded=True, author_id=31
        )
        _, rows = self.csv(section='summary')
        values = {row['Indicateur']: row['Valeur'] for row in rows}
        self.assertEqual(values['Rendez-vous tenu'], '1')
        self.assertEqual(
            values['Ventes ou contrats conclus sur rendez-vous tenus non annulés'], '1'
        )
        self.row.status = 'cancelled'
        self.row.save()
        _, rows = self.csv(section='summary')
        values = {row['Indicateur']: row['Valeur'] for row in rows}
        self.assertEqual(values['Rendez-vous tenu'], '0')
        self.assertEqual(values['Annulé'], '1')
        _, appointments = self.csv(section='appointments')
        self.assertEqual(appointments[0]['Statut du rendez-vous'], 'Annulé')
        self.assertEqual(appointments[0]['Présence déclarée dans le compte rendu'], 'Oui')
        self.assertEqual(appointments[0]['Vente ou contrat déclaré dans le compte rendu'], 'Oui')

    def test_appointment_and_declared_call_exports_identify_agents_without_dispatch(self):
        AppointmentReport.objects.create(
            appointment=self.row,
            performed=False,
            outcome='absent',
            note='Synthetic report',
            author_id=31,
        )
        _, rows = self.csv(section='appointments', status='absent')
        self.assertEqual(rows[0]['Auteur réservation ID'], '99')
        self.assertEqual(rows[0]['Auteur compte rendu ID'], '31')
        self.assertEqual(rows[0]['Statut du rendez-vous'], 'Absent')
        self.assertTrue(rows[0]['Début (UTC ISO 8601)'].endswith('+00:00'))
        ContactCallOutcome.objects.create(
            contact=self.row.contact,
            setter_id=21,
            request_id=uuid4(),
            fingerprint='x',
            result='voicemail',
            note='Synthetic voicemail',
        )
        _, calls = self.csv(section='calls')
        self.assertEqual(calls[0]['Auteur ID'], '21')
        self.assertEqual(calls[0]['Résultat déclaré'], 'Répondeur')

    def test_excessive_exports_never_silently_truncate_and_invalid_exports_are_not_audited(self):
        with patch('packages.exports.MAX_ROWS', {'csv': 1, 'xlsx': 1}):
            result = self.client.get(self.url, {'file': 'csv', 'section': 'prospects'})
            self.assertEqual(result.status_code, 400)
        for params in (
            {'file': 'zip'},
            {'file': 'csv', 'page': '1'},
            {'file': 'csv', 'section': 'secret'},
            {'file': 'csv', 'section': 'summary', 'search': 'x'},
            {'file': 'csv', 'section': 'calls', 'status': 'absent'},
        ):
            self.assertEqual(self.client.get(self.url, params).status_code, 400)
        self.assertFalse(BusinessAudit.objects.filter(action='client_folder.exported').exists())

    def test_exports_are_reserved_to_the_correct_workspace_administrator(self):
        for role in ('setter', 'advisor'):
            self.fixture.auth(31, role)
            self.assertEqual(self.client.get(self.url, {'file': 'csv'}).status_code, 403)
        self.client.credentials()
        self.assertIn(self.client.get(self.url, {'file': 'csv'}).status_code, (401, 403))
