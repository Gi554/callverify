import csv
from io import StringIO, BytesIO
from uuid import uuid4
from django.test import TestCase, override_settings
from rest_framework.test import APIClient
from packages.access.tickets import issue
from openpyxl import load_workbook
from .test_isolation_client import client_fixture
from .models import Prospect, IsolationVisit, IsolationVisitResult, IsolationFollowUp, AuditEvent


@override_settings(DEBUG=True, SECURE_SSL_REDIRECT=False)
class IsolationFolderExportTests(TestCase):
    def setUp(self):
        self.prospect, self.visit = client_fixture()
        self.client = APIClient()
        self.client.credentials(
            HTTP_AUTHORIZATION='Bearer ' + issue(99, 'isolation', 'administrator')
        )
        self.url = '/api/client-folders/72/export/'

    def test_csv_and_excel_exclude_other_clients_and_include_original_setter(self):
        other = Prospect.objects.create(
            campaign=self.prospect.campaign, raw_phone='secret', name='Foreign secret'
        )
        values = {
            f.attname: getattr(self.visit, f.attname)
            for f in IsolationVisit._meta.fields
            if not f.primary_key and f.name not in ('created_at', 'request_id')
        }
        IsolationVisit.objects.create(
            **{
                **values,
                'request_id': uuid4(),
                'prospect': other,
                'prospect_id': other.pk,
                'client_id': 73,
            }
        )
        result = self.client.get(self.url, {'file': 'csv', 'section': 'appointments'})
        self.assertEqual(result.status_code, 200)
        rows = list(csv.DictReader(StringIO(result.content.decode('utf-8-sig')), delimiter=';'))
        self.assertEqual(len(rows), 1)
        self.assertNotIn('Foreign secret', str(rows))
        self.assertEqual(rows[0]['Setter ID'], '21')
        self.assertEqual(rows[0]['Auteur réservation ID'], '91')
        book = self.client.get(self.url, {'file': 'xlsx', 'section': 'prospects'})
        self.assertEqual(book.status_code, 200)
        self.assertEqual(load_workbook(BytesIO(book.content)).active.max_row, 2)
        self.assertEqual(
            AuditEvent.objects.filter(action='isolation.client_folder.exported').count(), 2
        )

    def test_summaries_and_tasks_exclude_cancelled_sales_and_keep_notes(self):
        IsolationVisitResult.objects.create(
            visit=self.visit,
            performed=True,
            closed=True,
            outcome='completed',
            note='Synthetic result',
        )
        self.visit.status = 'cancelled'
        self.visit.save()
        summary = self.client.get(self.url, {'file': 'csv', 'section': 'summary'})
        self.assertEqual(summary.status_code, 200)
        rows = list(csv.DictReader(StringIO(summary.content.decode('utf-8-sig')), delimiter=';'))
        values = {row['Indicateur']: row['Valeur'] for row in rows}
        self.assertEqual(
            values['Ventes ou contrats conclus sur rendez-vous tenus non annulés'], '0'
        )
        IsolationFollowUp.objects.create(
            prospect=self.prospect,
            request_id=uuid4(),
            fingerprint='x',
            created_by_id=21,
            note='Synthetic task',
            due_at=self.visit.starts_at,
        )
        tasks = self.client.get(self.url, {'file': 'csv', 'section': 'tasks'})
        self.assertEqual(tasks.status_code, 200)
        self.assertIn('Synthetic task', tasks.content.decode('utf-8-sig'))

    def test_other_roles_and_another_space_cannot_export(self):
        for workspace, role in [
            ('isolation', 'setter'),
            ('isolation', 'client'),
            ('isolation', 'closer'),
            ('assurance', 'administrator'),
        ]:
            self.client.credentials(HTTP_AUTHORIZATION='Bearer ' + issue(99, workspace, role))
            self.assertIn(self.client.get(self.url, {'file': 'csv'}).status_code, (401, 403))
