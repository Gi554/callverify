import csv
import io
import re
from itertools import chain
import zipfile
from decimal import Decimal
from datetime import timedelta
import phonenumbers
from openpyxl import load_workbook, Workbook
from django.conf import settings
from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from twilio.rest import Client
from twilio.http.http_client import TwilioHttpClient
from .models import Campaign, PhoneNumber, Prospect, ImportBatch, ImportRow, LookupCheck, ExportJob, PlatformSettings

def normalize(value):
    # Excel numeric phone cells may arrive as an integral float or decimal text.
    text = str(value).strip()
    match = re.fullmatch(r'(\+?\d{10,15})\.0+', text)
    if match:
        text = match.group(1)
    try:
        number = phonenumbers.parse(text, 'CA')
        if not phonenumbers.is_valid_number(number):
            return None, '', 'Numéro invalide selon libphonenumber'
        country = phonenumbers.region_code_for_number(number) or ''
        if country != 'CA':
            return None, country, 'Numéro hors du périmètre canadien'
        return phonenumbers.format_number(number, phonenumbers.PhoneNumberFormat.E164), country, 'Validation locale libphonenumber'
    except phonenumbers.NumberParseException:
        return None, '', 'Format non reconnu'

def read_table(batch):
    with batch.original.open('rb') as stream:
        data = stream.read(settings.MAX_IMPORT_BYTES + 1)
    if len(data) > settings.MAX_IMPORT_BYTES:
        raise ValueError('Fichier trop volumineux')
    if batch.original.name.lower().endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Archive XLSX décompressée trop volumineuse')
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            worksheet = workbook[batch.sheet] if batch.sheet else workbook.active
            iterator = worksheet.iter_rows(values_only=True)
            first_row = next(iterator)
            if first_row and normalize(first_row[0])[0]:
                # A headerless phone list must keep its first number as data.
                headers = ['phone'] + [f'column_{i+1}' for i in range(1,len(first_row))]
                iterator = chain([first_row], iterator)
            else:
                headers = [str(v or '').strip() for v in first_row]
            rows = []
            for values in iterator:
                if len(rows) >= settings.MAX_IMPORT_ROWS:
                    raise ValueError('Nombre maximal de lignes dépassé')
                rows.append(dict(zip(headers, [str(v) if v is not None else '' for v in values])))
            sheets = workbook.sheetnames
        finally:
            workbook.close()
    else:
        try:
            content = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            content = data.decode('cp1252')
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(content), dialect=dialect)
        headers = reader.fieldnames or []
        rows = []
        for row in reader:
            if len(rows) >= settings.MAX_IMPORT_ROWS:
                raise ValueError('Nombre maximal de lignes dépassé')
            if None in row:
                raise ValueError('Nombre de colonnes incohérent')
            rows.append({key: value or '' for key, value in row.items()})
        sheets = []
    if not headers or any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError('En-têtes manquants ou dupliqués')
    return headers, rows, sheets

@transaction.atomic
def process_import(batch_id):
    batch = ImportBatch.objects.select_for_update().get(pk=batch_id)
    if batch.status == 'completed':
        return batch.report
    headers, rows, _ = read_table(batch)
    phone_column = batch.mapping.get('phone')
    if phone_column not in headers:
        raise ValueError('La colonne téléphone est obligatoire')
    if any(column not in headers for column in batch.mapping.values()):
        raise ValueError('Colonne du mappage inconnue')
    # Campaign lock serializes concurrent imports and preserves one prospect per number.
    Campaign.objects.select_for_update().get(pk=batch.campaign_id)
    report = {'total':len(rows), 'valid':0, 'invalid':0, 'duplicates':0, 'empty':0}
    for line, row in enumerate(rows, start=2):
        if not any(str(value).strip() for value in row.values()):
            report['empty'] += 1
            ImportRow.objects.create(batch=batch, line=line, source=row, error='Ligne vide')
            continue
        raw = row.get(phone_column, '').strip()
        e164, country, reason = normalize(raw)
        phone = None
        duplicate = False
        prospect = None
        if e164:
            phone, _ = PhoneNumber.objects.get_or_create(e164=e164, defaults={'country':country})
            prospect = Prospect.objects.filter(campaign=batch.campaign, phone=phone).first()
            duplicate = prospect is not None
        if prospect is None:
            prospect = Prospect.objects.create(campaign=batch.campaign, phone=phone, raw_phone=raw[:250],
                name=row.get(batch.mapping.get('name',''), '')[:250],
                reference=row.get(batch.mapping.get('reference',''), '')[:250],
                validity='valid' if e164 else 'invalid', reason=reason)
        ImportRow.objects.create(batch=batch, prospect=prospect, line=line, source=row,
            duplicate=duplicate, error='' if e164 else reason)
        report['duplicates' if duplicate else 'valid' if e164 else 'invalid'] += 1
    batch.status = 'completed'
    batch.report = report
    batch.save(update_fields=['status','report'])
    return report

def run_lookup(check_id):
    from .rate_limit import acquire_lookup
    acquire_lookup()
    with transaction.atomic():
        check = LookupCheck.objects.select_for_update().select_related('phone').get(pk=check_id)
        if check.status != 'pending':
            return
        check.status = 'running'
        check.save(update_fields=['status'])
    try:
        if not settings.LOOKUP_ENABLED:
            raise ValueError('Twilio Lookup désactivé')
        resource = Client(settings.TWILIO_ACCOUNT_SID, settings.TWILIO_AUTH_TOKEN,
            http_client=TwilioHttpClient(timeout=15)).lookups.v2.phone_numbers(check.phone.e164)
        result = resource.fetch(fields=check.packages) if check.packages else resource.fetch()
        payload = {'phone_number':result.phone_number, 'valid':result.valid,
            'validation_errors':result.validation_errors, 'country_code':result.country_code}
        check.response = payload
        check.status = 'completed'
        check.checked_at=timezone.now()
        update={'validity':'valid' if result.valid else 'invalid','line_status':'unknown'}
        package_errors=[]
        for field in filter(None,check.packages.split(',')):
            data=getattr(result,field,None)
            payload[field]=data if isinstance(data,dict) else None
            if isinstance(data,dict) and data.get('error_code'):
                package_errors.append(str(data['error_code']))
            elif isinstance(data,dict) and field=='line_type_intelligence':
                update.update(carrier=str(data.get('carrier_name') or '')[:200],line_type=str(data.get('type') or '')[:40])
            elif isinstance(data,dict) and field=='line_status':
                # Active is not proof of reachability; keep the raw provider status in history.
                value=str(data.get('status','')).lower()
                update['line_status']=value if value in ('reachable','unreachable') else 'unknown'
        if package_errors:
            check.error='Package indisponible: '+', '.join(package_errors)
            check.provider_code=','.join(package_errors)[:40]
        check.response=payload
        PhoneNumber.objects.filter(pk=check.phone_id).update(**update)
        Prospect.objects.filter(phone=check.phone,campaign=check.campaign).update(validity='valid' if result.valid else 'invalid',reason='Twilio Lookup Basic')
    except Exception as error:
        check.status = 'error'
        # Never store a provider exception string that could include credentials or URLs.
        check.error = 'Lookup indisponible: ' + type(error).__name__
        code=getattr(error,'code',None)
        check.provider_code=str(code)[:40] if isinstance(code,int) else ''
    check.save(update_fields=['response','status','error','provider_code','checked_at'])
    if not LookupCheck.objects.filter(campaign=check.campaign,status__in=['pending','running']).exists():
        Campaign.objects.filter(pk=check.campaign_id,status='active').update(status='completed')

def schedule_lookup(campaign, force=False):
    from .delivery import enqueue
    if not settings.LOOKUP_ENABLED or not settings.TWILIO_ACCOUNT_SID or not settings.TWILIO_AUTH_TOKEN:
        raise ValueError('Configurer et activer Twilio Lookup avant lancement')
    ids = []
    cached = 0
    with transaction.atomic():
        campaign = Campaign.objects.select_for_update().get(pk=campaign.pk)
        if campaign.status == 'archived':
            raise ValueError('Campagne archivée')
        phones = PhoneNumber.objects.filter(prospect__campaign=campaign).distinct()
        if not phones.exists():
            raise ValueError('Aucun numéro importé au format valide : consultez le rapport d’import avant de lancer Lookup')
        config,_=PlatformSettings.objects.get_or_create(pk=1)
        unit_cost = Decimal(settings.LOOKUP_UNIT_COST)
        if unit_cost < 0:
            raise ValueError('Coût unitaire invalide')
        for phone in phones:
            if LookupCheck.objects.filter(campaign=campaign, phone=phone, status__in=['pending','running']).exists():
                continue
            cutoff=timezone.now()-timedelta(hours=config.lookup_cache_hours)
            cached_check=LookupCheck.objects.filter(phone=phone,status='completed',packages=settings.LOOKUP_PACKAGES,error='').filter(
                Q(checked_at__gte=cutoff)|Q(checked_at__isnull=True,created_at__gte=cutoff)).order_by('-created_at').first() if not force else None
            if cached_check:
                cached += 1
                if not LookupCheck.objects.filter(campaign=campaign,phone=phone,status='completed',checked_at=cached_check.checked_at or cached_check.created_at).exists():
                    LookupCheck.objects.create(campaign=campaign,phone=phone,status='completed',packages=cached_check.packages,
                        checked_at=cached_check.checked_at or cached_check.created_at,response={**cached_check.response,'cache_source_check_id':cached_check.pk})
                Prospect.objects.filter(campaign=campaign,phone=phone).update(validity=phone.validity,reason='Twilio Lookup (cache)')
                continue
            if campaign.spent + unit_cost > campaign.budget:
                raise ValueError('Budget insuffisant pour le lot')
            check = LookupCheck.objects.create(campaign=campaign, phone=phone, cost=unit_cost,packages=settings.LOOKUP_PACKAGES)
            campaign.spent += unit_cost
            ids.append(check.pk)
        if phones.exists():campaign.status='active' if ids else 'completed'
        campaign.save(update_fields=['spent','status'])
        for check_id in ids:
            enqueue('lookup',check_id)
    return {'scheduled':len(ids), 'cached':cached, 'cost_type':'reserved_estimate'}

EXPORT_COLUMNS = ['name','reference','raw_phone','e164','validity','line_status','carrier','line_type','compliance','compliance_reason','reason','checked_at']
EXPORT_COLUMNS += ['lookup_classification','lookup_line_status','lookup_error','lookup_checked_at','voice_status','voice_rang','voice_answered','voice_amd','voice_error_code','voice_checked_at']
EXPORT_COLUMNS += ['line_state','line_error','line_checked_at']
FILTERS = {'validity', 'compliance', 'phone__line_status'}

def filtered_prospects(campaign, filters):
    if set(filters) - FILTERS:
        raise ValueError('Filtre inconnu')
    return Prospect.objects.filter(campaign=campaign, **filters).select_related('phone').order_by('pk')

def safe_cell(value):
    value = str(value or '')
    return "'" + value if value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n')) else value

@transaction.atomic
def build_export(job_id):
    job = ExportJob.objects.select_for_update().get(pk=job_id)
    if job.status == 'completed':
        return
    if job.expires_at<=timezone.now():
        job.status='expired';job.save(update_fields=['status']);return
    try:
        columns = job.columns or EXPORT_COLUMNS
        if set(columns) - set(EXPORT_COLUMNS):
            raise ValueError('Colonne inconnue')
        rows = []
        for prospect in filtered_prospects(job.campaign, job.filters):
            check = prospect.phone.checks.filter(status='completed',campaign=job.campaign).order_by('-created_at').first() if prospect.phone else None
            values = {column:getattr(prospect,column,'') for column in columns}
            values.update(e164=prospect.phone.e164 if prospect.phone else '',
                line_status=prospect.phone.line_status if prospect.phone else 'unknown',
                carrier=prospect.phone.carrier if prospect.phone else '',line_type=prospect.phone.line_type if prospect.phone else '',
                checked_at=(check.checked_at or check.created_at).isoformat() if check else '')
            from .qualification import technical_results
            values.update(technical_results(prospect))
            rows.append([str(values.get(column, '')) if values.get(column) is not None else '' for column in columns])
        if job.format == 'xlsx':
            book = Workbook()
            sheet = book.active
            sheet.title = 'Résultats'
            sheet.append(columns)
            for row in rows:
                sheet.append(row)
                for cell in sheet[sheet.max_row]:
                    cell.data_type='s'
                    cell.quotePrefix=True
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            output = io.BytesIO()
            book.save(output)
            content = output.getvalue()
        else:
            output = io.StringIO(newline='')
            writer = csv.writer(output)
            writer.writerow(columns)
            writer.writerows([[safe_cell(value) for value in row] for row in rows])
            content = output.getvalue().encode('utf-8-sig')
        job.file.save(f'results-{job.pk}.{job.format}', ContentFile(content), save=False)
        job.status = 'completed'
    except Exception as error:
        job.status = 'error'
        job.error = type(error).__name__
    job.save()
