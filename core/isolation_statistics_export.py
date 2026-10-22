"""Readable, scoped statistics exports; never include contact details or free notes."""

import csv
import io
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

LABELS = {
    'visits': 'Visites réservées',
    'prospects': 'Prospects distincts',
    'cancelled': 'Annulées',
    'upcoming': 'À venir',
    'reported': 'Résultats renseignés',
    'performed': 'Effectuées (Oui)',
    'not_performed': 'Non effectuées (Non)',
    'results_due': 'Résultats manquants',
    'interested': 'Intéressés',
    'not_interested': 'Non intéressés',
    'follow_up': 'Suivi nécessaire',
    'performed_rate': 'Taux effectué / résultats (%)',
    'interest_rate': 'Taux intérêt / effectuées (%)',
}
LABELS.update(
    closed='Ventes conclues déclarées',
    not_closed='Tenues : vente non conclue déclarée',
    closure_unknown='Tenues : vente non renseignée',
    unassigned_closer='Visites sans closer',
    closing_rate='Taux conclusion / tenues (%)',
)


def safe(value):
    if not isinstance(value, str):
        return value
    # Spreadsheet programs may treat leading whitespace before = as a formula.
    return (
        "'" + value
        if value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n'))
        else value
    )


def response(result, groups, kind):
    group = result.get('group', 'setter')
    person = 'Closer' if group == 'closer' else 'Setter'
    zone = result['timezone']
    meta = [str(result['from']), str(result['to']), result['as_of'].isoformat(), zone]
    headers = [
        'Périmètre',
        person + ' ID',
        person,
        *LABELS.values(),
        f'Du ({zone})',
        f'Au inclus ({zone})',
        'Calculé le (UTC)',
        'Fuseau de la période',
        'Définition',
    ]

    def values(label, row):
        return [
            label,
            row.get(group + '_id', ''),
            row.get(group + '_name', ''),
            *[row[key] if row[key] is not None else 'Non calculable' for key in LABELS],
            *meta,
            result['definition'],
        ]

    rows = [values('Total autorisé', result['totals']), *[values(person, row) for row in groups]]
    filename = f"statistiques-isolation-{result['from']}-{result['to']}.{kind}"
    if kind == 'csv':
        stream = io.StringIO(newline='')
        stream.write('\ufeff')
        writer = csv.writer(stream, delimiter=';')
        writer.writerow(headers)
        writer.writerows([[safe(value) for value in row] for row in rows])
        content = stream.getvalue().encode('utf-8')
        content_type = 'text/csv; charset=utf-8'
    else:
        book = Workbook()
        summary = book.active
        summary.title = 'Synthèse'
        summary.append(['Statistiques Isolation', 'Valeur'])
        for label, value in [
            (f'Du ({zone})', meta[0]),
            (f'Au inclus ({zone})', meta[1]),
            ('Calculé le (UTC)', meta[2]),
            ('Définition', result['definition']),
        ]:
            summary.append([label, value])
        for key, label in LABELS.items():
            summary.append(
                [
                    label,
                    result['totals'][key]
                    if result['totals'][key] is not None
                    else 'Non calculable',
                ]
            )
        summary.column_dimensions['A'].width = 44
        summary.column_dimensions['B'].width = 75
        summary['B5'].alignment = Alignment(wrap_text=True, vertical='top')
        summary.row_dimensions[5].height = 160
        if groups:
            sheet = book.create_sheet('Par ' + group)
            sheet.append(headers)
            for row in rows[1:]:
                sheet.append(row)
            sheet.auto_filter.ref = sheet.dimensions
            for i in range(1, len(headers) + 1):
                sheet.column_dimensions[get_column_letter(i)].width = 26
        for sheet in book:
            sheet.freeze_panes = 'A2'
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='53661C')
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'  # Force text, even malicious names beginning '='.
        stream = io.BytesIO()
        book.save(stream)
        content = stream.getvalue()
        content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    reply = HttpResponse(content, content_type=content_type)
    reply['Content-Disposition'] = f'attachment; filename="{filename}"'
    reply['Cache-Control'] = 'no-store'
    reply['X-Content-Type-Options'] = 'nosniff'
    return reply
