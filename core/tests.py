import io
import tempfile
from pathlib import Path
from datetime import timedelta
from decimal import Decimal
from unittest.mock import patch
from django.test import TestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from rest_framework.test import APIClient
from openpyxl import Workbook, load_workbook
from twilio.request_validator import RequestValidator
from django.contrib.auth.tokens import default_token_generator
from django.utils.http import urlsafe_base64_encode
from django.utils.encoding import force_bytes
from .models import *
from .services import normalize, process_import, build_export, schedule_lookup, run_lookup, safe_cell

@override_settings(CELERY_TASK_ALWAYS_EAGER=True, SECURE_SSL_REDIRECT=False,
    MFA_ENCRYPTION_KEY='MDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDA=',
    PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class BackendTests(TestCase):
    def setUp(self):
        from django.core.cache import cache
        cache.clear()
        test_root = Path(__file__).resolve().parent.parent / '.test-media'
        test_root.mkdir(exist_ok=True)
        self.storage = tempfile.TemporaryDirectory(dir=test_root)
        self.override = override_settings(MEDIA_ROOT=self.storage.name)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.addCleanup(self.storage.cleanup)
        self.admin = AdminUser.objects.create_user('admin', password='test-password-long', is_staff=True)
        self.other = AdminUser.objects.create_user('other', password='test-password-long', is_staff=True)
        self.client = APIClient()
        self.client.force_authenticate(self.admin)
        self.campaign = Campaign.objects.create(owner=self.admin,name='Test',purpose='Qualification',source='Démo',budget=1)
    def batch(self, content):
        return ImportBatch.objects.create(campaign=self.campaign,
            original=SimpleUploadedFile('test.csv',content.encode()),mapping={'phone':'phone','name':'name'})
    def imported(self):
        batch = self.batch('phone,name\n4165550123,Alice\n')
        process_import(batch.pk)
        return Prospect.objects.get(campaign=self.campaign)
    def test_normalization(self):
        self.assertEqual(normalize('(416) 555-0123')[0], '+14165550123')
        self.assertIsNone(normalize('abc')[0])
        self.assertIsNone(normalize('+12025550123')[0])
    def test_import_deduplication_and_invalids(self):
        batch = self.batch('phone,name\n4165550123,A\n+14165550123,B\nabc,C\n')
        report = process_import(batch.pk)
        self.assertEqual(report['duplicates'],1)
        self.assertEqual(report['invalid'],1)
        self.assertEqual(ImportRow.objects.count(),3)
        self.assertEqual(Prospect.objects.count(),2)
        self.assertEqual(process_import(batch.pk),report)
        self.assertEqual(ImportRow.objects.count(),3)
    def test_duplicate_across_batches(self):
        self.imported()
        batch = self.batch('phone,name\n4165550123,B\n')
        self.assertEqual(process_import(batch.pk)['duplicates'],1)
        self.assertEqual(Prospect.objects.count(),1)
    def test_xlsx_import(self):
        workbook = Workbook()
        workbook.active.append(['phone','name'])
        workbook.active.append([4165550123,'Alice'])
        output = io.BytesIO()
        workbook.save(output)
        batch = ImportBatch.objects.create(campaign=self.campaign,
            original=SimpleUploadedFile('test.xlsx',output.getvalue()),mapping={'phone':'phone'})
        self.assertEqual(process_import(batch.pk)['valid'],1)
    def test_import_rollback_invalid_mapping(self):
        batch = self.batch('other,name\n4165550123,A\n')
        with self.assertRaises(ValueError):
            process_import(batch.pk)
        self.assertEqual(ImportRow.objects.count(),0)
    def test_owner_boundary(self):
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(f'/api/campaigns/{self.campaign.pk}/').status_code,404)
        self.assertEqual(self.client.get('/api/campaigns/').data['count'],0)
        response = self.client.post('/api/imports/',{'campaign':self.campaign.pk,'file':SimpleUploadedFile('test.csv',b'phone\n4165550123')})
        self.assertEqual(response.status_code,404)
    def test_employee_denied(self):
        employee = AdminUser.objects.create_user('employee')
        self.client.force_authenticate(employee)
        self.assertEqual(self.client.get('/api/campaigns/').status_code,403)
    def test_anonymous_denied(self):
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get('/api/campaigns/').status_code,403)
    def test_sensitive_endpoints_denied(self):
        self.assertEqual(self.client.get('/api/administrators/').status_code,403)
        self.assertEqual(self.client.get('/api/audit/').status_code,403)
    def test_import_preview_confirm(self):
        response = self.client.post('/api/imports/', {'campaign':self.campaign.pk,
            'file':SimpleUploadedFile('sample.csv',b'phone,name\n4165550123,Alice\n')})
        self.assertEqual(response.status_code,201)
        batch_id = response.data['id']
        with self.captureOnCommitCallbacks(execute=True):
            result = self.client.post(f'/api/imports/{batch_id}/confirm/',{'mapping':{'phone':'phone','name':'name'}},format='json')
        self.assertEqual(result.status_code,202)
        self.assertEqual(ImportBatch.objects.get(pk=batch_id).status,'completed')
        self.assertEqual(self.client.post(f'/api/imports/{batch_id}/confirm/',{},format='json').status_code,400)
    def test_export_formulas_and_filters(self):
        prospect = self.imported()
        prospect.name = '=HYPERLINK("bad")'
        prospect.save()
        job = ExportJob.objects.create(campaign=self.campaign,author=self.admin,columns=['name','e164'],
            filters={'validity':'valid'},expires_at=timezone.now()+timedelta(hours=1))
        build_export(job.pk)
        job.refresh_from_db()
        with job.file.open('rb') as stream:
            self.assertIn(b"'=HYPERLINK",stream.read())
        response = self.client.get(f'/api/exports/{job.pk}/download/')
        self.assertEqual(response.status_code,200)
        # Release only the download handles: response.close() also emits
        # request_finished, which closes PostgreSQL's enclosing test transaction.
        for closer in response._resource_closers:
            closer()
        job.expires_at = timezone.now()-timedelta(seconds=1)
        job.save()
        self.assertEqual(self.client.get(f'/api/exports/{job.pk}/download/').status_code,400)
    def test_xlsx_export_is_text(self):
        prospect = self.imported()
        prospect.name = '=SUM(A1)'
        prospect.save()
        job = ExportJob.objects.create(campaign=self.campaign,author=self.admin,format='xlsx',columns=['name'],expires_at=timezone.now()+timedelta(hours=1))
        build_export(job.pk)
        job.refresh_from_db()
        with job.file.open('rb') as stream:
            book = load_workbook(io.BytesIO(stream.read()))
        self.assertEqual(book.active['A2'].value,'=SUM(A1)')
        self.assertEqual(book.active['A2'].data_type,'s')
    def test_formula_whitespace(self):
        self.assertTrue(safe_cell('  =1+1').startswith("'"))
    @override_settings(LOOKUP_ENABLED=True,TWILIO_ACCOUNT_SID='test',TWILIO_AUTH_TOKEN='test',LOOKUP_UNIT_COST='2')
    def test_budget_is_atomic(self):
        self.imported()
        with self.assertRaises(ValueError):
            schedule_lookup(self.campaign)
        self.assertEqual(LookupCheck.objects.count(),0)
        self.campaign.refresh_from_db()
        self.assertEqual(self.campaign.spent,0)
    def test_lookup_disabled(self):
        self.assertEqual(self.client.post(f'/api/campaigns/{self.campaign.pk}/lookup/',{},format='json').status_code,400)
    @override_settings(LOOKUP_ENABLED=True,TWILIO_ACCOUNT_SID='test',TWILIO_AUTH_TOKEN='test')
    @patch('core.services.Client')
    def test_lookup_does_not_imply_reachability(self, client):
        prospect = self.imported()
        result = client.return_value.lookups.v2.phone_numbers.return_value.fetch.return_value
        result.phone_number = prospect.phone.e164
        result.valid = True
        result.validation_errors = []
        result.country_code = 'CA'
        check = LookupCheck.objects.create(campaign=self.campaign,phone=prospect.phone)
        run_lookup(check.pk)
        prospect.phone.refresh_from_db()
        check.refresh_from_db()
        self.assertEqual(check.status,'completed')
        self.assertEqual(prospect.phone.line_status,'unknown')
    @override_settings(LOOKUP_ENABLED=True,TWILIO_ACCOUNT_SID='test',TWILIO_AUTH_TOKEN='test')
    @patch('core.services.Client',side_effect=RuntimeError('secret'))
    def test_lookup_error_no_secrets(self, client):
        prospect = self.imported()
        check = LookupCheck.objects.create(campaign=self.campaign,phone=prospect.phone)
        run_lookup(check.pk)
        check.refresh_from_db()
        self.assertEqual(check.status,'error')
        self.assertNotIn('secret',check.error)
    def test_call_start_blocked(self):
        campaign = CallCampaign.objects.create(campaign=self.campaign)
        self.assertEqual(self.client.post(f'/api/calls/{campaign.pk}/start/').status_code,409)
    @override_settings(TWILIO_AUTH_TOKEN='test-secret')
    def test_signed_webhook_idempotent_and_monotonic(self):
        prospect = self.imported()
        campaign = CallCampaign.objects.create(campaign=self.campaign)
        attempt = CallAttempt.objects.create(call_campaign=campaign,prospect=prospect,sid='CA123')
        url = '/api/webhooks/twilio/voice/'
        payload = {'CallSid':'CA123','CallStatus':'completed','SequenceNumber':'3','CallDuration':'20'}
        signature = RequestValidator('test-secret').compute_signature('http://testserver'+url,payload)
        for _ in range(2):
            self.assertEqual(self.client.post(url,payload,HTTP_X_TWILIO_SIGNATURE=signature).status_code,204)
        self.assertEqual(WebhookEvent.objects.count(),1)
        payload.update(CallStatus='ringing',SequenceNumber='1')
        signature = RequestValidator('test-secret').compute_signature('http://testserver'+url,payload)
        self.client.post(url,payload,HTTP_X_TWILIO_SIGNATURE=signature)
        attempt.refresh_from_db()
        self.assertEqual(attempt.status,'completed')
    @override_settings(TWILIO_AUTH_TOKEN='test-secret')
    def test_unsigned_webhook_denied(self):
        self.assertEqual(self.client.post('/api/webhooks/twilio/voice/',{'CallSid':'x'}).status_code,403)
    def test_login_csrf_enforced(self):
        client = APIClient(enforce_csrf_checks=True)
        self.assertEqual(client.post('/api/auth/login/',{'username':'admin','password':'test-password-long'}).status_code,403)
    def test_frontend_proxy_login_csrf(self):
        client = APIClient(enforce_csrf_checks=True)
        token = client.get('/api/auth/csrf/').data['csrfToken']
        result = client.post('/api/auth/login/',{'username':'admin','password':'test-password-long'},
            HTTP_X_CSRFTOKEN=token,HTTP_ORIGIN='http://127.0.0.1:5173')
        self.assertEqual(result.status_code,200)
    def test_superadmin_requires_mfa(self):
        AdminUser.objects.create_superuser('super','super@example.com','test-password-long')
        client = APIClient()
        self.assertEqual(client.post('/api/auth/login/',{'username':'super','password':'test-password-long'}).status_code,401)
    def test_reset_is_single_use(self):
        token = default_token_generator.make_token(self.admin)
        payload = {'uid':urlsafe_base64_encode(force_bytes(self.admin.pk)), 'token':token,
            'password':'Fresh-password-9347!'}
        client = APIClient()
        self.assertEqual(client.post('/api/auth/reset/confirm/',payload).status_code,200)
        self.assertEqual(client.post('/api/auth/reset/confirm/',payload).status_code,400)
        self.admin.refresh_from_db()
        self.assertTrue(self.admin.check_password(payload['password']))
    def test_invalid_reset_token(self):
        self.assertEqual(APIClient().post('/api/auth/reset/confirm/',{'uid':'bogus','token':'bogus'}).status_code,400)
    @override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
    def test_recovery_generic_response(self):
        from django.core import mail
        self.admin.email = 'admin@example.com'
        self.admin.save()
        client = APIClient()
        known = client.post('/api/auth/reset/request/',{'email':'admin@example.com'})
        unknown = client.post('/api/auth/reset/request/',{'email':'missing@example.com'})
        self.assertEqual(known.data,unknown.data)
        self.assertEqual(len(mail.outbox),1)
    def test_invitation_activation(self):
        superadmin = AdminUser.objects.create_superuser('root','root@example.com','password')
        self.client.force_authenticate(superadmin)
        response = self.client.post('/api/administrators/',{'username':'invited','email':'invited@example.com','invite':True},format='json')
        self.assertEqual(response.status_code,201)
        user = AdminUser.objects.get(username='invited')
        self.assertFalse(user.has_usable_password())
        payload = {**response.data['invitation'],'password':'Fresh-password-9847!'}
        self.assertEqual(APIClient().post('/api/auth/reset/confirm/',payload).status_code,200)
        user.refresh_from_db()
        self.assertTrue(user.has_usable_password())
    @patch('core.views.totp_step',return_value=100)
    def test_mfa_code_not_reusable(self, step):
        AdminUser.objects.create_superuser('root','root@example.com','test-password-long',mfa_secret='test')
        client = APIClient()
        payload = {'username':'root','password':'test-password-long','code':'123456'}
        self.assertEqual(client.post('/api/auth/login/',payload).status_code,200)
        client.post('/api/auth/session/')
        self.assertEqual(client.post('/api/auth/login/',payload).status_code,401)
    def test_other_admin_cannot_download_export(self):
        job = ExportJob.objects.create(campaign=self.campaign,author=self.admin,expires_at=timezone.now()+timedelta(hours=1))
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(f'/api/exports/{job.pk}/download/').status_code,404)
    @override_settings(LOOKUP_ENABLED=True,TWILIO_ACCOUNT_SID='test',TWILIO_AUTH_TOKEN='test')
    def test_cached_lookup_not_charged(self):
        prospect = self.imported()
        LookupCheck.objects.create(campaign=self.campaign,phone=prospect.phone,status='completed')
        result = schedule_lookup(self.campaign)
        self.assertEqual(result['cached'],1)
        self.assertEqual(result['scheduled'],0)
        self.campaign.refresh_from_db()
        self.assertEqual(self.campaign.spent,0)
    def test_invalid_export_filters_rejected(self):
        result = self.client.post('/api/exports/',{'campaign':self.campaign.pk,'filters':{'campaign__owner':'2'}},format='json')
        self.assertEqual(result.status_code,400)
    def test_mfa_secret_encrypted(self):
        from .security import decrypt_secret
        self.admin.mfa_secret='JBSWY3DPEHPK3PXP';self.admin.save()
        self.admin.refresh_from_db()
        self.assertTrue(self.admin.mfa_secret.startswith('enc:'))
        self.assertEqual(decrypt_secret(self.admin.mfa_secret),'JBSWY3DPEHPK3PXP')
    def test_backup_code_single_use(self):
        from .security import generate_backup_codes
        self.admin.mfa_enabled=True;self.admin.save()
        codes=generate_backup_codes(self.admin)
        client=APIClient()
        payload={'username':'admin','password':'test-password-long','code':codes[0]}
        self.assertEqual(client.post('/api/auth/login/',payload).status_code,200)
        client.post('/api/auth/session/')
        self.assertEqual(client.post('/api/auth/login/',payload).status_code,401)
    @patch('core.tasks.execute_delivery.delay',side_effect=ConnectionError('offline'))
    def test_outbox_persists_when_broker_unavailable(self,task):
        from .delivery import enqueue
        with self.captureOnCommitCallbacks(execute=True):
            delivery=enqueue('import',self.batch('phone,name\n4165550123,A\n').pk)
        delivery.refresh_from_db()
        self.assertEqual(delivery.status,'pending')
        self.assertEqual(delivery.attempts,1)
        self.assertEqual(delivery.error,'ConnectionError')
    def test_task_delivery_idempotent(self):
        from .delivery import enqueue
        from .tasks import execute_delivery
        batch=self.batch('phone,name\n4165550123,A\n')
        with self.captureOnCommitCallbacks(execute=True):delivery=enqueue('import',batch.pk)
        execute_delivery(delivery.pk)
        self.assertEqual(ImportRow.objects.count(),1)
    def test_platform_settings_superadmin_only(self):
        self.assertEqual(self.client.get('/api/settings/').status_code,403)
        root=AdminUser.objects.create_superuser('root','root@example.com','password')
        self.client.force_authenticate(root)
        self.assertEqual(self.client.patch('/api/settings/',{'lookup_rate':0},format='json').status_code,400)
        self.assertEqual(self.client.patch('/api/settings/',{'lookup_rate':3},format='json').status_code,200)
        self.assertEqual(PlatformSettings.objects.get(pk=1).lookup_rate,3)
    def test_compliance_requires_valid_evidence(self):
        prospect=self.imported()
        url=f'/api/prospects/{prospect.pk}/classify/'
        payload={'compliance':'allowed','reason':'Preuve examinée par administrateur'}
        self.assertEqual(self.client.post(url,payload,format='json').status_code,400)
        ConsentEvidence.objects.create(prospect=prospect,basis='Accord explicite',evidence='Preuve fictive',recorded_by=self.admin,expires_at=timezone.now()+timedelta(days=1))
        self.assertEqual(self.client.post(url,payload,format='json').status_code,200)
        Suppression.objects.create(e164=prospect.phone.e164,reason='Opposition')
        self.assertEqual(self.client.post(url,payload,format='json').status_code,400)
    def test_call_eligibility_rechecks_exclusions(self):
        prospect=self.imported();prospect.compliance='allowed';prospect.save()
        ConsentEvidence.objects.create(prospect=prospect,basis='Accord',evidence='Preuve',recorded_by=self.admin,expires_at=timezone.now()+timedelta(days=1))
        campaign=CallCampaign.objects.create(campaign=self.campaign)
        url=f'/api/calls/{campaign.pk}/eligibility/'
        self.assertEqual(self.client.get(url).data['eligible'],1)
        Suppression.objects.create(e164=prospect.phone.e164,reason='Opposition')
        self.assertEqual(self.client.get(url).data['eligible'],0)
    def test_report_private_and_formula_safe(self):
        batch=self.batch('phone,name\nabc,A\n')
        process_import(batch.pk)
        response=self.client.get(f'/api/imports/{batch.pk}/report/')
        self.assertEqual(response.status_code,200)
        self.assertIn(b'ligne',response.content)
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(f'/api/imports/{batch.pk}/report/').status_code,404)
    def test_export_purge_keeps_history(self):
        from .maintenance import purge_files
        self.imported()
        job=ExportJob.objects.create(campaign=self.campaign,author=self.admin,expires_at=timezone.now()+timedelta(hours=1))
        build_export(job.pk);job.refresh_from_db();name=job.file.name;storage=job.file.storage
        job.expires_at=timezone.now()-timedelta(seconds=1);job.save()
        self.assertEqual(purge_files(),1)
        job.refresh_from_db();self.assertEqual(job.status,'expired');self.assertFalse(storage.exists(name))
    def test_result_search_and_pagination(self):
        self.imported()
        result=self.client.get('/api/prospects/?search=Alice&page_size=1')
        self.assertEqual(result.data['count'],1)
        self.assertEqual(len(result.data['results']),1)
        self.assertEqual(self.client.get('/api/prospects/?search=missing').data['count'],0)
    def test_notifications_scoped(self):
        batch=self.batch('phone,name\nabc,A\n');batch.status='error';batch.save()
        self.assertEqual(len(self.client.get('/api/notifications/').data['items']),1)
        self.client.force_authenticate(self.other)
        self.assertEqual(len(self.client.get('/api/notifications/').data['items']),0)
    def test_health(self):
        self.assertEqual(APIClient().get('/api/health/').status_code,200)
    @override_settings(CELERY_TASK_ALWAYS_EAGER=False)
    @patch('redis.Redis.from_url')
    def test_global_rate_limiter_blocks(self,redis):
        from .rate_limit import acquire_lookup,RateLimited
        redis.return_value.eval.return_value=0
        with self.assertRaises(RateLimited):acquire_lookup()
    def test_mfa_enrollment_generates_qr_and_backup(self):
        response=self.client.post('/api/auth/mfa/',{'action':'enroll'},format='json')
        self.assertEqual(response.status_code,200)
        self.assertTrue(response.data['qr'].startswith('data:image/svg+xml;base64,'))
        with patch('core.views.valid_totp',return_value=True):
            response=self.client.post('/api/auth/mfa/',{'code':'123456'},format='json')
        self.assertEqual(len(response.data['backup_codes']),10)
    @override_settings(LOOKUP_ENABLED=True,TWILIO_ACCOUNT_SID='test',TWILIO_AUTH_TOKEN='test')
    @patch('core.services.Client')
    def test_advanced_active_is_not_reachable(self,client):
        prospect=self.imported()
        result=client.return_value.lookups.v2.phone_numbers.return_value.fetch.return_value
        result.phone_number=prospect.phone.e164;result.valid=True;result.validation_errors=[];result.country_code='CA'
        result.line_status={'status':'Active','error_code':None}
        check=LookupCheck.objects.create(campaign=self.campaign,phone=prospect.phone,packages='line_status')
        run_lookup(check.pk);prospect.phone.refresh_from_db();check.refresh_from_db()
        self.assertEqual(prospect.phone.line_status,'unknown')
        self.assertEqual(check.response['line_status']['status'],'Active')
    @override_settings(LOOKUP_ENABLED=True,TWILIO_ACCOUNT_SID='test',TWILIO_AUTH_TOKEN='test')
    @patch('core.services.Client')
    def test_package_restriction_keeps_basic_result(self,client):
        prospect=self.imported()
        result=client.return_value.lookups.v2.phone_numbers.return_value.fetch.return_value
        result.phone_number=prospect.phone.e164;result.valid=True;result.validation_errors=[];result.country_code='CA'
        result.line_type_intelligence={'error_code':60601}
        check=LookupCheck.objects.create(campaign=self.campaign,phone=prospect.phone,packages='line_type_intelligence')
        run_lookup(check.pk);check.refresh_from_db()
        self.assertEqual(check.status,'completed');self.assertEqual(check.provider_code,'60601')
    def test_superadmin_promotion_requires_mfa(self):
        root=AdminUser.objects.create_superuser('root','root@example.com','password')
        self.client.force_authenticate(root)
        payload={'id':self.other.pk,'role':'superadministrator'}
        self.assertEqual(self.client.patch('/api/administrators/',payload,format='json').status_code,400)
        self.other.mfa_enabled=True;self.other.save()
        self.assertEqual(self.client.patch('/api/administrators/',payload,format='json').status_code,200)
    def test_invoice_reconciliation_requires_superadmin(self):
        prospect=self.imported();check=LookupCheck.objects.create(campaign=self.campaign,phone=prospect.phone)
        url=f'/api/lookups/{check.pk}/reconcile/'
        body={'actual_cost':'0.01','invoice_reference':'FACTURE-TEST'}
        self.assertEqual(self.client.post(url,body,format='json').status_code,403)
        root=AdminUser.objects.create_superuser('root','root@example.com','password');self.client.force_authenticate(root)
        self.assertEqual(self.client.post(url,body,format='json').status_code,200)
        check.refresh_from_db();self.assertEqual(check.actual_cost,Decimal('0.01'))
    def test_dashboard_deduplicates_across_campaigns(self):
        prospect=self.imported()
        other=Campaign.objects.create(owner=self.admin,name='Deuxième campagne',purpose='Test',source='Fictif')
        Prospect.objects.create(campaign=other,phone=prospect.phone,raw_phone=prospect.raw_phone,validity='valid')
        result=self.client.get('/api/dashboard/').data
        self.assertEqual(result['unique_numbers'],1)
        self.assertEqual(next(r['count'] for r in result['validity'] if r['validity']=='valid'),1)


    def voice_observation(self, amd=False):
        prospect=self.imported()
        obj=CallCampaign.objects.create(campaign=self.campaign)
        attempt=CallAttempt.objects.create(call_campaign=obj,prospect=prospect,sid='CA-observation',amd_requested=amd)
        return prospect, attempt
    def post_voice_event(self, payload, amd=False):
        url='/api/webhooks/twilio/amd/' if amd else '/api/webhooks/twilio/voice/'
        signature=RequestValidator('test-secret').compute_signature('http://testserver'+url,payload)
        return self.client.post(url,payload,HTTP_X_TWILIO_SIGNATURE=signature)
    @override_settings(VOICE_ENABLED=True)
    def test_voice_start_disabled_even_when_overridden(self):
        obj=CallCampaign.objects.create(campaign=self.campaign)
        self.assertEqual(self.client.post(f'/api/calls/{obj.pk}/start/',{'confirm_test':True,'prospects':[1]},format='json').status_code,409)
    @patch('twilio.rest.Client')
    def test_old_voice_delivery_cannot_emit(self, provider):
        from .voice import submit_attempt
        _, attempt=self.voice_observation()
        attempt.status='pending';attempt.save()
        submit_attempt(attempt.pk)
        provider.assert_not_called()
        attempt.refresh_from_db();self.assertEqual(attempt.status,'blocked')
    @override_settings(TWILIO_AUTH_TOKEN='test-secret')
    def test_ringing_and_answered_are_distinct(self):
        prospect,attempt=self.voice_observation()
        self.assertEqual(self.post_voice_event({'CallSid':attempt.sid,'CallStatus':'ringing','SequenceNumber':'1'}).status_code,204)
        attempt.refresh_from_db();self.assertIsNotNone(attempt.rang_at);self.assertIsNone(attempt.answered_at)
        self.post_voice_event({'CallSid':attempt.sid,'CallStatus':'completed','SequenceNumber':'2'})
        attempt.refresh_from_db();self.assertIsNotNone(attempt.answered_at)
        self.assertEqual(attempt.amd,'unknown')
        from .qualification import technical_results
        self.assertEqual(technical_results(prospect)['voice_amd'],'not_requested')
    @override_settings(TWILIO_AUTH_TOKEN='test-secret')
    def test_async_amd_without_call_status_is_idempotent(self):
        prospect,attempt=self.voice_observation(amd=True)
        payload={'CallSid':attempt.sid,'AnsweredBy':'machine_start'}
        self.assertEqual(self.post_voice_event(payload,amd=True).status_code,204)
        self.post_voice_event(payload,amd=True)
        attempt.refresh_from_db();self.assertEqual(attempt.amd,'machine_start')
        self.assertEqual(attempt.status,'queued')
        self.assertEqual(WebhookEvent.objects.count(),1)
        self.assertIsNone(attempt.answered_at)
    @override_settings(TWILIO_AUTH_TOKEN='test-secret')
    def test_voice_provider_error_and_history_are_saved(self):
        prospect,attempt=self.voice_observation()
        self.post_voice_event({'CallSid':attempt.sid,'CallStatus':'failed','ErrorCode':'21217'})
        attempt.refresh_from_db();self.assertEqual(attempt.error_code,'21217')
        result=self.client.get(f'/api/prospects/{prospect.pk}/history/').data
        self.assertEqual(result['voice'][0]['status'],'failed')
        self.assertEqual(len(result['voice_events']),1)
    def test_lookup_active_does_not_imply_reachability(self):
        from .qualification import technical_results
        prospect=self.imported()
        check=LookupCheck.objects.create(campaign=self.campaign,phone=prospect.phone,status='completed',response={'valid':True,'line_status':{'status':'Active'}})
        result=technical_results(prospect)
        self.assertEqual(result['lookup_classification'],'valid')
        self.assertFalse(result['voice_answered'])
        check.response={'valid':True,'line_status':{'status':'reachable'}};check.save()
        self.assertEqual(technical_results(prospect)['lookup_classification'],'potentially_reachable')
    def test_export_keeps_lookup_and_voice_separate(self):
        prospect,attempt=self.voice_observation()
        attempt.status='busy';attempt.save()
        job=ExportJob.objects.create(campaign=self.campaign,author=self.admin,format='csv',
            columns=['lookup_classification','voice_status','voice_rang','voice_answered','voice_amd'],expires_at=timezone.now()+timedelta(hours=1))
        build_export(job.pk);job.refresh_from_db()
        with job.file.open('rb') as stream: content=stream.read().decode('utf-8-sig')
        self.assertIn('not_checked,busy,False,False,not_requested',content)
    def test_amd_options_are_optional(self):
        from .voice import amd_parameters
        self.assertEqual(amd_parameters(False,''),{})
        self.assertEqual(amd_parameters(True,'https://example.invalid/api/webhooks/twilio/amd/')['machine_detection'],'Enable')

    def test_excel_integral_decimal_phone_is_normalized(self):
        self.assertEqual(normalize('15144325221.0')[0], '+15144325221')
        self.assertEqual(normalize(15144325221.0)[0], '+15144325221')
        self.assertIsNone(normalize('15144325221.5')[0])
    def test_headerless_excel_keeps_first_phone(self):
        book=Workbook();sheet=book.active
        sheet.append([15813079131.0]);sheet.append([15144325221.0])
        data=io.BytesIO();book.save(data)
        batch=ImportBatch.objects.create(campaign=self.campaign,original=SimpleUploadedFile('phones.xlsx',data.getvalue()),mapping={'phone':'phone'})
        report=process_import(batch.pk)
        self.assertEqual(report['total'],2)
        self.assertEqual(report['valid'],2)
    @override_settings(LOOKUP_ENABLED=True,TWILIO_ACCOUNT_SID='test',TWILIO_AUTH_TOKEN='test')
    def test_empty_campaign_cannot_schedule_lookup(self):
        with self.assertRaisesMessage(ValueError,'Aucun numéro importé'):
            schedule_lookup(self.campaign)

    def test_line_state_is_independent_of_newer_basic_validation(self):
        from .qualification import technical_results
        prospect=self.imported()
        LookupCheck.objects.create(campaign=self.campaign,phone=prospect.phone,status='completed',packages='line_status',response={'valid':True,'line_status':{'status':'inactive','error_code':None}})
        LookupCheck.objects.create(campaign=self.campaign,phone=prospect.phone,status='completed',packages='',response={'valid':True})
        result=technical_results(prospect)
        self.assertEqual(result['line_state'],'inactive')
        self.assertEqual(result['lookup_classification'],'valid')
        dashboard=self.client.get('/api/dashboard/').data
        self.assertEqual(next(r['count'] for r in dashboard['line_states'] if r['state']=='inactive'),1)
    def test_basic_valid_and_60600_do_not_confirm_line(self):
        from .qualification import technical_results
        prospect=self.imported()
        LookupCheck.objects.create(campaign=self.campaign,phone=prospect.phone,status='completed',packages='',response={'valid':True})
        self.assertEqual(technical_results(prospect)['line_state'],'unknown')
        LookupCheck.objects.create(campaign=self.campaign,phone=prospect.phone,status='completed',packages='line_status',error='Package indisponible: 60600',response={'valid':True,'line_status':{'status':None,'error_code':60600}})
        self.assertEqual(technical_results(prospect)['line_state'],'unknown')

    @override_settings(DEBUG=True)
    def test_isolation_rejects_assurance_ticket(self):
        from packages.access.tickets import issue
        client=APIClient();client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(1,'assurance','administrator'))
        self.assertIn(client.get('/api/dashboard/').status_code,[401,403])
    @override_settings(DEBUG=True)
    def test_isolation_rejects_setter_ticket_on_admin_routes(self):
        from packages.access.tickets import issue
        client=APIClient();client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(1,'isolation','setter'))
        self.assertIn(client.get('/api/campaigns/').status_code,[401,403])
