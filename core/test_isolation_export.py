import csv
from io import BytesIO, StringIO
from openpyxl import load_workbook
from django.test import TestCase, override_settings
from rest_framework.test import APIClient
from packages.access.tickets import issue
from .models import Prospect, AuditEvent
from .test_isolation_work import fixture


@override_settings(DEBUG=True, SECURE_SSL_REDIRECT=False)
class IsolationExportTests(TestCase):
    def setUp(self):
        self.row = fixture()
        self.row.assigned_setter_id = 21
        self.row.assigned_setter_name = 'setter.demo01'
        self.row.reference = 'ISO-1'
        self.row.save()
        self.other = Prospect.objects.create(
            campaign=self.row.campaign, name='-1+cmd', raw_phone='+1 (416) 555-0199'
        )
        self.client = APIClient()
        self.url = '/api/work-queue/export/'

    def auth(self, pk, role, space='isolation'):
        self.client.credentials(HTTP_AUTHORIZATION='Bearer ' + issue(pk, space, role))

    def csv_rows(self, query=None):
        response = self.client.get(self.url, {'file': 'csv', **(query or {})})
        self.assertEqual(response.status_code, 200)
        self.assertIn('prospects-isolation-', response['Content-Disposition'])
        return list(csv.reader(StringIO(response.content.decode('utf-8')[1:]), delimiter=';'))

    def test_setter_exports_only_assigned_prospects(self):
        self.auth(21, 'setter')
        rows = self.csv_rows()
        self.assertEqual(len(rows), 2)
        self.assertEqual(
            rows[1][:8],
            [str(self.row.pk), 'ISO-1', 'Prospect', '4165550123']
            + [
                '+14165550123',
                'Isolation',
                'setter.demo01',
                'À contacter',
            ],
        )
        audit = AuditEvent.objects.get(action='isolation.prospects.exported')
        self.assertEqual(audit.metadata['actor_portal_id'], 21)
        self.assertEqual(audit.metadata['rows'], 1)

    def test_administrator_exports_team_with_filters(self):
        self.auth(91, 'administrator')
        rows = self.csv_rows()
        self.assertEqual([row[0] for row in rows[1:]], [str(self.other.pk), str(self.row.pk)])
        self.assertEqual(rows[1][2], "'-1+cmd")
        self.assertEqual(rows[1][3], '+1 (416) 555-0199')
        self.assertEqual(rows[1][6], 'Non affecté')
        self.assertEqual(len(self.csv_rows({'owner': 'unassigned'})), 2)
        response = self.client.get(self.url, {'file': 'xlsx', 'owner': '21'})
        sheet = load_workbook(BytesIO(response.content)).worksheets[0]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['C2'].value, 'Prospect')

    def test_invalid_requests_are_refused(self):
        self.auth(91, 'administrator')
        self.assertEqual(self.client.get(self.url).status_code, 400)
        self.assertEqual(self.client.get(self.url, {'file': 'csv', 'owner': 'x'}).status_code, 400)
        self.auth(5, 'client')
        self.assertEqual(self.client.get(self.url, {'file': 'csv'}).status_code, 403)
