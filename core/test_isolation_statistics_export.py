import csv
import io
from django.test import TestCase,override_settings
from rest_framework.test import APIClient
from openpyxl import load_workbook
from packages.access.tickets import issue
from . import test_isolation_statistics as fixtures
from .models import AuditEvent,IsolationVisit,Prospect

@override_settings(DEBUG=True,SECURE_SSL_REDIRECT=False)
class IsolationStatisticsExportTests(TestCase):
    def setUp(self):
        self.fixture=fixtures.IsolationStatisticsTests(methodName='test_empty_rates_are_not_false_zero');self.fixture.setUp()
        self.client=self.fixture.client;self.url='/api/workspace-statistics/'
    def auth(self,pk,role,space='isolation'):self.client.credentials(HTTP_AUTHORIZATION='Bearer '+issue(pk,space,role))
    def export(self,kind='csv',**params):return self.client.get(self.url,{'export':kind,**params})
    def rows(self,response):return list(csv.reader(io.StringIO(response.content.decode('utf-8-sig')),delimiter=';'))
    def test_csv_readability_metadata_and_audit(self):
        self.fixture.visit(performed=True,outcome='interested')
        response=self.export();self.assertEqual(response.status_code,200)
        self.assertTrue(response.content.startswith(b'\xef\xbb\xbf'))
        rows=self.rows(response);self.assertIn('Visites réservées',rows[0]);self.assertEqual(rows[1][0],'Total autorisé')
        self.assertIn('Définition',rows[0]);self.assertIn('Taux effectué',rows[1][-1])
        self.assertEqual(response['Cache-Control'],'no-store');self.assertEqual(response['X-Content-Type-Options'],'nosniff')
        self.assertIn('attachment',response['Content-Disposition'])
        audit=AuditEvent.objects.get(action='isolation.statistics.exported')
        self.assertEqual(audit.metadata['actor_portal_id'],91);self.assertEqual(audit.metadata['scope'],'administrator')
        self.assertNotIn('Confidentiel',response.content.decode('utf-8-sig'))
    def test_exports_all_groups_ignoring_page_and_no_formula_execution(self):
        for index in range(27):self.fixture.visit(setter=100+index)
        IsolationVisit.objects.filter(setter_id=100).update(setter_name='=HYPERLINK("https://invalid")')
        csv_rows=self.rows(self.export(page=999))
        self.assertEqual(len(csv_rows),29);self.assertTrue(csv_rows[2][2].startswith("'="))
        response=self.export('xlsx',page=999);self.assertEqual(response.status_code,200)
        book=load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames,['Synthèse','Par setter']);self.assertEqual(book['Par setter'].max_row,28)
        self.assertEqual(book['Par setter']['C2'].data_type,'s');self.assertTrue(book['Par setter']['C2'].value.startswith('='))
        self.assertEqual(book['Par setter'].freeze_panes,'A2')
        self.assertTrue(book['Par setter'].auto_filter.ref)
    def test_client_export_scope_and_names_do_not_leak(self):
        self.fixture.visit(performed=True)
        foreign=self.fixture.visit(setter=22,client=73);foreign.setter_name='SECRET';foreign.save()
        self.auth(72,'client')
        for kind in ('csv','xlsx'):
            response=self.export(kind);self.assertEqual(response.status_code,200)
            if kind=='csv':self.assertNotIn('SECRET',response.content.decode('utf-8-sig'));self.assertEqual(len(self.rows(response)),3)
            else:
                book=load_workbook(io.BytesIO(response.content));self.assertEqual(book['Par setter'].max_row,2)
                self.assertNotIn('SECRET',str(list(book['Par setter'].values)))
    def test_setter_summary_only_and_transferred_prospect_removed(self):
        self.fixture.visit();self.fixture.visit(setter=22)
        self.auth(21,'setter');rows=self.rows(self.export());self.assertEqual(len(rows),2)
        self.assertEqual(rows[1][1:3],['',''])
        book=load_workbook(io.BytesIO(self.export('xlsx').content));self.assertEqual(book.sheetnames,['Synthèse'])
        Prospect.objects.filter(pk=self.fixture.prospect.pk).update(assigned_setter_id=22)
        rows=self.rows(self.export());self.assertEqual(rows[1][3],'0')
    def test_dates_empty_rates_and_permissions(self):
        self.fixture.visit()
        day=self.fixture.day.date().isoformat()
        rows=self.rows(self.export(**{'from':day,'to':day}));self.assertEqual(rows[1][3],'1')
        self.assertIn('Non calculable',rows[1])
        self.assertEqual(self.export(**{'from':'bad'}).status_code,400)
        self.assertEqual(self.export('pdf').status_code,400)
        self.auth(21,'setter','assurance');self.assertEqual(self.export().status_code,401)
        self.client.credentials();self.assertEqual(self.export('xlsx').status_code,401)
