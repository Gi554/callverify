import csv
import io
from datetime import datetime,timedelta,timezone as utc
from django.test import TestCase,override_settings
from openpyxl import load_workbook
from . import test_isolation_statistics as fixtures
from .isolation_statistics import period
from .models import IsolationVisitResult,IsolationVisit,AuditEvent,Prospect


@override_settings(DEBUG=True,SECURE_SSL_REDIRECT=False)
class PerformanceTests(TestCase):
    setUp=fixtures.IsolationStatisticsTests.setUp
    auth=fixtures.IsolationStatisticsTests.auth
    visit=fixtures.IsolationStatisticsTests.visit

    def report(self,performed=None,closed=None,closer=77,setter=21,**changes):
        row=self.visit(performed=performed,closer_id=closer,closer_name=f'Closer {closer}',setter=setter,**changes)
        if performed is not None:IsolationVisitResult.objects.filter(visit=row).update(closed=closed)
        return row

    def test_twenty_assigned_fifteen_held_five_sales(self):
        for i in range(20):self.report(performed=i<15,closed=i<5)
        data=self.client.get(self.url,{'group':'closer'}).data
        totals=data['totals'];self.assertEqual(totals['visits'],20)
        self.assertEqual((totals['performed'],totals['not_performed'],totals['closed']),(15,5,5))
        self.assertEqual(totals['performed_rate'],75)
        self.assertEqual(totals['closing_rate'],33.33)
        self.assertEqual(data['closers']['count'],1)
        self.assertEqual(data['closers']['results'][0]['closing_rate'],33.33)

    def test_pending_cancelled_and_legacy_unknown_are_not_false_outcomes(self):
        self.report(performed=True,closed=True)
        self.report(performed=True,closed=False)
        self.report(performed=True,closed=None)
        self.report(performed=False,closed=False)
        self.report();self.report(status='cancelled')
        data=self.client.get(self.url).data['totals']
        self.assertEqual((data['closed'],data['not_closed'],data['closure_unknown']),(1,1,1))
        self.assertEqual(data['results_due'],1);self.assertEqual(data['not_performed'],1)
        self.assertEqual(data['performed_rate'],75);self.assertEqual(data['closing_rate'],33.33)
        IsolationVisitResult.objects.all().delete()
        empty=self.client.get(self.url).data['totals']
        self.assertIsNone(empty['closing_rate']);self.assertIsNone(empty['performed_rate'])

    def test_week_boundaries_in_canadian_timezone_and_dst(self):
        start,end,lower,upper,zone=period({'week_of':'2026-11-01','timezone':'America/Toronto'},datetime.now(utc.utc))
        self.assertEqual(str(start),'2026-10-26');self.assertEqual(str(end),'2026-11-01')
        self.assertEqual((upper-lower).total_seconds()/3600,169)
        spring=period({'week_of':'2027-03-14','timezone':'America/Toronto'},datetime.now(utc.utc))
        self.assertEqual((spring[3]-spring[2]).total_seconds()/3600,167)
        for instant in (lower-timedelta(seconds=1),lower,upper-timedelta(seconds=1),upper):
            row=self.report();row.starts_at=instant;row.ends_at=instant+timedelta(minutes=30);row.save()
        response=self.client.get(self.url,{'week_of':'2026-11-01','timezone':'America/Toronto'})
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.data['totals']['visits'],2)
        self.assertEqual(str(response.data['from']),'2026-10-26')

    def test_closer_scope_applies_before_totals_groups_and_exports(self):
        self.report(performed=True,closed=True)
        foreign=self.report(performed=False,closer=88,setter=22);foreign.closer_name='SECRET';foreign.save()
        self.auth(77,'closer')
        for group in ('setter','closer'):
            response=self.client.get(self.url,{'group':group})
            self.assertEqual(response.status_code,200,response.data)
            self.assertEqual(response.data['totals']['visits'],1)
            self.assertNotIn('closers',response.data);self.assertNotIn('setters',response.data)
            exported=self.client.get(self.url,{'group':group,'export':'csv'})
            rows=list(csv.reader(io.StringIO(exported.content.decode('utf-8-sig')),delimiter=';'))
            self.assertEqual(len(rows),2);self.assertNotIn('SECRET',str(rows))
        IsolationVisit.objects.filter(closer_id=77).update(closer_id=88)
        self.assertEqual(self.client.get(self.url).data['totals']['visits'],0)

    def test_admin_and_client_group_current_closer_and_keep_unassigned(self):
        own=self.report(performed=True,closed=True)
        self.report(closer=None)
        foreign=self.report(client=73,closer=88);foreign.closer_name='SECRET';foreign.save()
        self.auth(72,'client')
        data=self.client.get(self.url,{'group':'closer'}).data
        self.assertEqual(data['closers']['count'],2)
        self.assertEqual(data['closers']['results'][0]['closer_name'],'Non attribué')
        self.assertNotIn('SECRET',str(data));self.assertEqual(data['totals']['unassigned_closer'],1)
        IsolationVisit.objects.filter(pk=own.pk).update(closer_id=99,closer_name='Nouveau')
        data=self.client.get(self.url,{'group':'closer'}).data
        self.assertEqual(data['closers']['results'][1]['closer_id'],99)
        self.assertEqual(data['closers']['results'][1]['closed'],1)

    def test_setter_scope_still_excludes_transferred_prospects(self):
        self.report(performed=True,closed=True);self.report(setter=22,closer=88)
        self.auth(21,'setter')
        data=self.client.get(self.url,{'group':'closer'}).data
        self.assertEqual(data['totals']['closed'],1);self.assertNotIn('closers',data)
        Prospect.objects.filter(pk=self.prospect.pk).update(assigned_setter_id=22)
        self.assertEqual(self.client.get(self.url).data['totals']['visits'],0)

    def test_closer_exports_include_all_groups_and_safe_names(self):
        for index in range(27):self.report(closer=100+index)
        IsolationVisit.objects.filter(closer_id=100).update(closer_name='=HYPERLINK("https://invalid")')
        response=self.client.get(self.url,{'export':'csv','group':'closer','page':999,'timezone':'America/Toronto'})
        rows=list(csv.reader(io.StringIO(response.content.decode('utf-8-sig')),delimiter=';'))
        self.assertEqual(len(rows),29);self.assertEqual(rows[0][1:3],['Closer ID','Closer'])
        self.assertTrue(rows[2][2].startswith("'="))
        self.assertIn('America/Toronto',rows[1])
        book=load_workbook(io.BytesIO(self.client.get(self.url,{'export':'xlsx','group':'closer'}).content))
        self.assertEqual(book.sheetnames,['Synthèse','Par closer'])
        self.assertEqual(book['Par closer'].max_row,28)
        self.assertEqual(book['Par closer']['C2'].data_type,'s')
        self.assertTrue(book['Par closer'].auto_filter.ref)
        self.assertNotIn('Confidentiel',str(list(book['Par closer'].values)))
        audit=AuditEvent.objects.filter(action='isolation.statistics.exported').first()
        self.assertEqual(audit.metadata['group'],'closer')

    def test_paginated_groups_do_not_change_totals(self):
        for index in range(27):self.report(closer=100+index)
        first=self.client.get(self.url,{'group':'closer'}).data
        self.assertEqual(first['closers']['count'],27);self.assertEqual(len(first['closers']['results']),25)
        second=self.client.get(self.url,{'group':'closer','page':2}).data
        self.assertEqual(len(second['closers']['results']),2)
        self.assertEqual(first['totals'],second['totals'])

    def test_invalid_and_conflicting_periods_are_rejected(self):
        for params in ({'timezone':'Mars/Test'},{'group':'client'},{'week_of':'bad'},
                       {'week_of':'2026-10-20','from':'2026-10-01'},
                       {'week_of':'9999-12-31'}, {'from':'0001-01-01','to':'0001-01-01','timezone':'Asia/Tokyo'}):
            response=self.client.get(self.url,params)
            self.assertEqual(response.status_code,400,(params,response.data))
